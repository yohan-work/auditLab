"""회사가 감사인에게 제출하는 원본 자료(더미)를 실무 양식 그대로 만든다.

산출물 (data/company/):
  01_합계잔액시산표_2025.xlsx          더존 Smart A 형식(추정)
  02_분개장_2025.xlsx                   더존 분개장 내보내기 형식(추정)
  03_분개장_2026_1-2월.xlsx             후속기간 분개장
  04_계정별원장_외상매출금_2025.xlsx    더존 계정별원장 형식(추정)
  05_거래처원장_외상매출금_2025.xlsx    거래처별 잔액(추정)
  06_거래처등록현황.xlsx                거래처 코드·사업자번호
  07_매출전자세금계산서목록_2025.xlsx   홈택스 목록조회 내려받기 형식(실측 기록 기준)
  08_매출전자세금계산서목록_2026_1-2월.xlsx
  09_은행거래내역_KB국민_보통예금.xlsx  KB국민은행 거래내역 엑셀 형식(공개 파서 기준)
  10_부가가치세신고서_과세표준요약_2025.xlsx
그리고 data/truth.json: 심어 둔 이상 거래 정답지 (엔진은 읽지 않음, 비교용)
"""
import json
import math
import random
from collections import defaultdict
from datetime import date, datetime, time, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from common import (BAND_FILL, BOLD, COMPANY, COMPANY_DIR, DATA, FY_END, FY_START, SUB_END,
                    body_row, d2s, header_row, month_end, next_biz, widths)

rng = random.Random(20251231)

# ---------------------------------------------------------------- 거래처
NAMES = [
    "한결산업", "대성테크", "미래산업", "태성기계", "동원정밀", "세진오토텍", "우림전자", "광명금속",
    "청운산업", "보람테크", "신화정공", "다온시스템", "해솔기계", "가람전기", "누리오토", "백산엔지니어링",
    "서해정밀", "한울산업", "제일모빌리티", "대한파워텍", "성우하이텍", "유진기공", "삼보테크", "금강정밀",
    "새론산업", "로뎀테크", "아림공업", "예성전자", "정우메탈", "하나기연", "도담산업", "온새미로",
    "푸른기계", "케이원텍", "에이스정공", "경인오토", "남부기계", "동해산업", "무궁화정밀", "비전테크",
]
PRODUCTS = [
    ("정밀가공부품", "A-100", 18_500), ("정밀가공부품", "A-220", 24_800), ("브래킷", "B-220", 6_400),
    ("샤프트", "S-15", 31_200), ("샤프트", "S-30", 44_700), ("하우징", "H-80", 52_000),
    ("기어세트", "G-12", 87_500), ("플랜지", "F-50", 12_300), ("금형부품", "M-07", 126_000),
]


def make_bizno(i):
    a = 100 + (i * 37) % 800
    mid = rng.choice(["81", "86", "87", "88"])
    return f"{a:03d}-{mid}-{rng.randint(10000, 99999):05d}"


CEOS = "김민수 이서연 박지훈 최유진 정하늘 강도윤 조수빈 윤재원 장서우 임하준".split()

customers = []
for i, n in enumerate(NAMES):
    weight = 1.0 / (i + 1) ** 0.85  # 상위 거래처 집중
    customers.append({
        "code": f"{101 + i:05d}",
        "name": f"(주){n}" if i % 3 else f"{n}(주)",
        "base": n,
        "bizno": make_bizno(i),
        "ceo": CEOS[i % len(CEOS)],
        "addr": rng.choice(["경기도 시흥시", "경기도 화성시", "인천광역시 남동구", "충청남도 아산시",
                            "경상남도 창원시", "경기도 평택시", "울산광역시 북구"]) + f" 가상로 {rng.randint(1, 300)}",
        "email": f"acct{101 + i}@customer.example",
        "lam": max(0.6, 7.5 * weight),
        "billing": "monthly" if i % 4 == 0 else "per",
        "terms": rng.choice([30, 45, 60, 60, 90]),
        "pay": rng.choice(["exact", "exact", "exact", "fee", "combined", "split"]),
        "payer": rng.choice([f"{n}(주)", f"(주){n}", f"{n}주식회사", f"주){n}", f"{n}주식회사"[:6]]),
    })
C = {c["code"]: c for c in customers}
by_base = {c["base"]: c for c in customers}

# ---------------------------------------------------------------- 매출(세금계산서·원장)
appr_seq = defaultdict(int)


def approval_no(issue):
    appr_seq[issue] += 1
    return f"{issue:%Y%m%d}-{41000000 + appr_seq[issue]:08d}-{rng.getrandbits(32):08x}"


def pick_items(target=None):
    lines = []
    for _ in range(rng.choice([1, 1, 2, 2, 3])):
        name, spec, price = rng.choice(PRODUCTS)
        qty = max(10, int(rng.lognormvariate(5.2, 0.7)))
        lines.append({"name": name, "spec": spec, "qty": qty, "price": price, "supply": qty * price})
    return lines


invoices = []      # 홈택스 목록 (작성일자 기준)
ledger_sales = []  # 원장 매출 전표 (외상매출금 차변)
open_ar = []       # 수금 대상 (청구 단위)


def add_invoice(cust, wdate, items, idate=None, kind="일반", orig=None, note=""):
    supply = sum(x["supply"] for x in items)
    vat = int(supply * 0.1) if supply >= 0 else -int(-supply * 0.1)
    idate = idate or wdate
    inv = {
        "appr": approval_no(idate), "wdate": wdate, "idate": idate, "sdate": next_biz(idate + timedelta(days=1)),
        "cust": cust["code"], "supply": supply, "vat": vat, "total": supply + vat,
        "kind": kind, "orig": orig, "items": items, "note": note,
    }
    invoices.append(inv)
    return inv


def add_ledger(cust, d, supply, vat, memo, inv=None, tag=None):
    e = {"date": d, "cust": cust["code"], "supply": supply, "vat": vat, "total": supply + vat,
         "memo": memo, "inv": inv["appr"] if inv else None, "tag": tag}
    ledger_sales.append(e)
    return e


def item_memo(items):
    first = items[0]
    return f"제품매출 {first['name']} {first['spec']}" + (f" 외 {len(items) - 1}건" if len(items) > 1 else "")


def months(y0, m0, y1, m1):
    y, m = y0, m0
    while (y, m) <= (y1, m1):
        yield y, m
        m += 1
        if m == 13:
            y, m = y + 1, 1


for c in customers:
    for y, m in months(2025, 1, 2026, 2):
        n = sum(1 for _ in range(30) if rng.random() < c["lam"] / 30)
        if n == 0:
            continue
        days = sorted(rng.sample([d for d in range(1, month_end(y, m).day + 1)
                                  if date(y, m, d).weekday() < 5 and (m, d) != (1, 1)], min(n, 18)))
        ships = [(date(y, m, d), pick_items()) for d in days]
        if c["billing"] == "monthly":
            items = []
            for sd, its in ships:
                for it in its:
                    items.append({**it, "date": sd})
            wd = month_end(y, m)
            idate = next_biz(wd + timedelta(days=rng.randint(1, 9)))
            inv = add_invoice(c, wd, items, idate=idate, note="월합계")
            add_ledger(c, wd, inv["supply"], inv["vat"], f"{m}월 제품매출(월합계)", inv)
        else:
            for sd, its in ships:
                for it in its:
                    it["date"] = sd
                inv = add_invoice(c, sd, its)
                add_ledger(c, sd, inv["supply"], inv["vat"], item_memo(its), inv)

# 전기이월 채권 (2024년 11~12월 매출분, 2025년 1~2월 회수)
opening = {}
for c in customers:
    amt = int(c["lam"] * 2 * 26_000_000 * rng.uniform(0.6, 1.2) / 1000) * 1000
    if amt > 0:
        opening[c["code"]] = amt

# ---------------------------------------------------------------- 이상 거래 심기
truth = []


def T(id_, kind, assertion, desc, **keys):
    truth.append({"id": id_, "type": kind, "assertion": assertion, "desc": desc, **keys})


def cust_of(base):
    return by_base[base]


def items_for(target):
    name, spec, price = rng.choice(PRODUCTS)
    qty = max(1, target // price)
    return [{"name": name, "spec": spec, "qty": qty, "price": price, "supply": qty * price}]


def ensure_inv(cust, months_, min_total=0):
    """해당 월에 그 거래처 일반 세금계산서가 있으면 가장 큰 것을, 없으면 하나 만들어 돌려준다."""
    cands = [i for i in invoices if i["cust"] == cust["code"] and i["kind"] == "일반" and i["wdate"].year == 2025
             and i["wdate"].month in months_ and i["total"] >= min_total]
    if cands:
        return max(cands, key=lambda i: i["total"])
    d = next_biz(date(2025, months_[0], 10))
    its = items_for(max(min_total, 11_000_000))
    for it in its:
        it["date"] = d
    inv = add_invoice(cust, d, its)
    add_ledger(cust, d, inv["supply"], inv["vat"], item_memo(its), inv)
    return inv


# A1 가공매출 의심: 신규 거래처, 원장에만 있고 세금계산서·입금 없음
ghost = {"code": "00141", "name": "(주)소망유통", "base": "소망유통", "bizno": "214-87-30561", "ceo": "오세훈",
         "addr": "서울특별시 금천구 가상로 77", "email": "som@customer.example", "lam": 0, "billing": "per",
         "terms": 60, "pay": "exact", "payer": "소망유통"}
customers.append(ghost)
C[ghost["code"]] = ghost
e = add_ledger(ghost, date(2025, 12, 15), 18_000_000, 1_800_000, "제품매출 정밀가공부품 A-220", tag="A1")
T("A1", "원장O·세금계산서X·입금X", "실재성", "기말 직전 신규 거래처 매출이 원장에만 있음. 세금계산서·입금 모두 없음 (가공매출 의심)",
  ledger_date=d2s(e["date"]), cust=ghost["code"], amount=e["total"])

# A2 세금계산서 미발행: 원장 매출·입금은 있으나 세금계산서 없음
c2 = cust_of("가람전기")
e = add_ledger(c2, date(2025, 6, 18), 7_300_000, 730_000, "제품매출 플랜지 F-50 외 1건", tag="A2")
T("A2", "원장O·세금계산서X·입금O", "발생(세무)", "매출·입금은 있으나 세금계산서 미발행 (부가세 신고 누락 위험)",
  ledger_date=d2s(e["date"]), cust=c2["code"], amount=e["total"])

# A3 매출 누락: 세금계산서·입금은 있는데 원장에 매출 없음 (입금은 가수금 처리)
c3 = cust_of("서해정밀")
inv = add_invoice(c3, date(2025, 9, 22), items_for(12_000_000))
T("A3", "세금계산서O·원장X", "완전성", "세금계산서 발행·입금 완료, 원장 매출 미계상 (입금은 가수금으로 처리)",
  appr=inv["appr"], cust=c3["code"], amount=inv["total"])
a3_inv = inv

# A4 금액 불일치(자릿수 전치): 원장 25,300,000 vs 세금계산서 23,500,000
c4 = cust_of("세진오토텍")
inv = add_invoice(c4, date(2025, 11, 12), [{"name": "하우징", "spec": "H-80", "qty": 500, "price": 47_000,
                                            "supply": 23_500_000, "date": date(2025, 11, 12)}])
add_ledger(c4, date(2025, 11, 12), 25_300_000, 2_530_000, "제품매출 하우징 H-80", inv, tag="A4")
T("A4", "금액 불일치", "정확성", "원장 공급가액 25,300,000 vs 세금계산서 23,500,000 (자릿수 전치 입력). 채권 1,980,000 과대",
  appr=inv["appr"], cust=c4["code"], amount=1_980_000)

# A5 반품 수정세금계산서(-) 원장 미반영
c5 = cust_of("우림전자")
base_inv = ensure_inv(c5, [9], min_total=9_000_000)
ret_items = [{"name": base_inv["items"][0]["name"], "spec": base_inv["items"][0]["spec"], "qty": -1,
              "price": 6_000_000, "supply": -6_000_000, "date": date(2025, 10, 14)}]
inv = add_invoice(c5, date(2025, 10, 14), ret_items, kind="수정", orig=base_inv["appr"],
                  note=f"환입 (당초승인번호 {base_inv['appr']})")
T("A5", "수정세금계산서 원장 미반영", "정확성·실재성", "10월 반품 수정세금계산서(-6,600,000)가 원장에 반영 안 됨. 채권 6,600,000 과대",
  appr=inv["appr"], orig=base_inv["appr"], cust=c5["code"], amount=6_600_000)
a5 = (base_inv, inv)

# A6 기간귀속(과대): 원장은 12/30 매출, 세금계산서 작성일·출고는 2026-01-05
c6 = cust_of("대성테크")
its = items_for(48_000_000)
for it in its:
    it["date"] = date(2026, 1, 5)
inv = add_invoice(c6, date(2026, 1, 5), its)
add_ledger(c6, date(2025, 12, 30), inv["supply"], inv["vat"], item_memo(its), inv, tag="A6")
T("A6", "기간귀속(당기 과대)", "기간귀속", "원장 2025-12-30 계상, 세금계산서 작성일·출고 2026-01-05. 2025년 매출 과대",
  appr=inv["appr"], cust=c6["code"], amount=inv["total"])

# A7 기간귀속(과소): 세금계산서 작성일 2025-12-29, 원장은 2026-01-05 계상
c7 = cust_of("성우하이텍")
its = items_for(9_000_000)
for it in its:
    it["date"] = date(2025, 12, 29)
inv = add_invoice(c7, date(2025, 12, 29), its)
add_ledger(c7, date(2026, 1, 5), inv["supply"], inv["vat"], item_memo(its), inv, tag="A7")
T("A7", "기간귀속(당기 과소)", "기간귀속·완전성", "세금계산서 작성일 2025-12-29, 원장은 2026-01-05 계상. 2025년 매출 과소",
  appr=inv["appr"], cust=c7["code"], amount=inv["total"])

# A8 기말 직전 대형 매출, 후속기간 전액 반품
c8 = cust_of("한결산업")
its = items_for(80_000_000)
for it in its:
    it["date"] = date(2025, 12, 26)
inv = add_invoice(c8, date(2025, 12, 26), its)
add_ledger(c8, date(2025, 12, 26), inv["supply"], inv["vat"], item_memo(its), inv, tag="A8")
ret = [{**its[0], "qty": -its[0]["qty"], "supply": -its[0]["supply"], "date": date(2026, 1, 16)}]
rinv = add_invoice(c8, date(2026, 1, 16), ret, kind="수정", orig=inv["appr"], note=f"환입 (당초승인번호 {inv['appr']})")
add_ledger(c8, date(2026, 1, 16), rinv["supply"], rinv["vat"], "매출환입 " + its[0]["name"], rinv, tag="A8r")
T("A8", "기말 직전 매출 후속 취소", "실재성·기간귀속", "2025-12-26 대형 매출이 2026-01-16 전액 반품(수정세금계산서). 밀어내기 매출 의심",
  appr=inv["appr"], ret_appr=rinv["appr"], cust=c8["code"], amount=inv["total"])
a8_inv = inv

# A9 제3자 입금: 매출처가 아닌 명의로 입금
c9 = cust_of("태성기계")
a9_inv = ensure_inv(c9, [7, 8], min_total=15_000_000)
T("A9", "제3자 명의 입금", "실재성(특수관계)", "태성기계 매출대금이 '태성홀딩스' 명의로 입금. 거래 상대방 확인 필요",
  appr=a9_inv["appr"], cust=c9["code"], amount=a9_inv["total"])

# A10 원장 중복 계상
c10 = cust_of("광명금속")
a10_inv = ensure_inv(c10, [11], min_total=12_000_000)
dup = next(e for e in ledger_sales if e["inv"] == a10_inv["appr"])
add_ledger(c10, dup["date"], dup["supply"], dup["vat"], dup["memo"], tag="A10")
T("A10", "원장 중복 계상", "실재성", "같은 세금계산서로 원장에 두 번 계상. 채권·매출 과대",
  appr=a10_inv["appr"], cust=c10["code"], amount=dup["total"])

# A11 장기 미회수 (평가): 3월·5월 청구분이 기말 후까지 미입금
c11 = cust_of("미래산업")
a11 = [ensure_inv(c11, [3]), ensure_inv(c11, [5])]
T("A11", "장기 미회수", "평가", "3월·5월 청구분이 기말 후 2개월까지 미입금 (대손충당금 검토 필요)",
  appr=[i["appr"] for i in a11], cust=c11["code"], amount=sum(i["total"] for i in a11))
a11_set = {i["appr"] for i in a11}

# ---------------------------------------------------------------- 수금(은행 입금)
deposits = []   # 은행 입금
ledger_coll = []  # 원장 수금 (외상매출금 대변)


def add_dep(d, amt, payer, kind="타행이체", links=(), tag=None):
    dt = datetime.combine(d, time(rng.randint(9, 16), rng.randint(0, 59), rng.randint(0, 59)))
    x = {"dt": dt, "amount": amt, "payer": payer, "kind": kind, "links": list(links), "tag": tag}
    deposits.append(x)
    return x


def book_coll(cust, d, ar_amt, fee=0, acct="108", memo="외상대금 회수"):
    ledger_coll.append({"date": d, "cust": cust, "ar": ar_amt, "fee": fee, "acct": acct, "memo": memo})


def due_of(c, wdate):
    return next_biz(month_end(wdate.year, wdate.month) + timedelta(days=c["terms"] + rng.randint(-3, 6)))


# 전기이월 회수
for code, amt in opening.items():
    c = C[code]
    d = next_biz(date(2025, 1, 1) + timedelta(days=rng.randint(10, 55)))
    add_dep(d, amt, c["payer"], links=[("OPEN", code)])
    book_coll(code, d, amt, memo="전기 외상대금 회수")

skip = {a8_inv["appr"], a5[1]["appr"]} | a11_set
skip |= {i["appr"] for i in invoices if i["wdate"] >= date(2026, 1, 1)}  # 후속기간 청구는 대부분 2월 이후 만기
by_cust_due = defaultdict(list)
for inv in invoices:
    if inv["appr"] in skip or inv["kind"] == "수정":
        continue
    c = C[inv["cust"]]
    amt = inv["total"]
    if inv is a5[0]:
        amt += a5[1]["total"]  # 반품분 차감 후 입금
    by_cust_due[(c["code"], due_of(c, inv["wdate"]))].append((inv, amt))

for (code, due), lst in sorted(by_cust_due.items(), key=lambda kv: kv[0][1]):
    c = C[code]
    if due > SUB_END:
        continue
    if c["pay"] == "combined" and len(lst) > 1:
        total = sum(a for _, a in lst)
        add_dep(due, total, c["payer"], links=[i["appr"] for i, _ in lst])
        for inv, a in lst:
            book_coll(code, due, a)
        continue
    for inv, amt in lst:
        payer = c["payer"]
        if inv is a9_inv:
            payer = "태성홀딩스"
        if c["pay"] == "split" and amt > 5_000_000:
            first = int(amt * 0.5 / 1000) * 1000
            add_dep(due, first, payer, links=[inv["appr"]])
            book_coll(code, due, first)
            d2 = next_biz(due + timedelta(days=30))
            if d2 <= SUB_END:
                add_dep(d2, amt - first, payer, links=[inv["appr"]])
                book_coll(code, d2, amt - first)
        elif c["pay"] == "fee":
            add_dep(due, amt - 1_000, payer, links=[inv["appr"]])
            book_coll(code, due, amt, fee=1_000)
        else:
            add_dep(due, amt, payer, links=[inv["appr"]])
            book_coll(code, due, amt)

# A2 입금 (세금계산서 없이)
add_dep(next_biz(date(2025, 8, 4)), 8_030_000, c2["payer"], links=["A2"], tag="A2")
book_coll(c2["code"], next_biz(date(2025, 8, 4)), 8_030_000)
# A3 입금은 원장에서 가수금 처리 (외상매출금 아님)
d = next_biz(date(2025, 11, 24))
add_dep(d, a3_inv["total"], c3["payer"], links=[a3_inv["appr"]], tag="A3")
book_coll(c3["code"], d, a3_inv["total"], acct="257", memo="입금 내역 확인중")

# A4: 고객은 세금계산서 금액만 입금 → 원장 수금은 그 금액, 차액 1,980,000 잔존 (기본 흐름에서 이미 처리됨)
# A10: 고객은 한 번만 입금 (기본 흐름)

# ---------------------------------------------------------------- 기타 은행 거래·전표 (시산표용)
journal = []  # (date, entries[(acct, cust, memo, dr, cr)])


def J(d, lines):
    journal.append((d, lines))


ACCT = {
    "103": "보통예금", "108": "외상매출금", "109": "대손충당금", "135": "부가세대급금", "153": "원재료",
    "206": "기계장치", "207": "감가상각누계액", "251": "외상매입금", "254": "예수금", "255": "부가세예수금",
    "257": "가수금", "331": "자본금", "375": "이월이익잉여금", "404": "제품매출", "455": "제품매출원가",
    "801": "급여", "811": "복리후생비", "818": "감가상각비", "831": "지급수수료", "901": "이자수익",
}
SUPPLIERS = ["(주)대한철강", "한국베어링(주)", "(주)동성소재", "삼일특수강(주)", "(주)경기공구", "남양화학(주)"]

withdrawals = []


def add_wd(d, amt, payee, kind="인터넷"):
    dt = datetime.combine(d, time(rng.randint(9, 17), rng.randint(0, 59), rng.randint(0, 59)))
    withdrawals.append({"dt": dt, "amount": amt, "payee": payee, "kind": kind})


for e in ledger_sales:
    J(e["date"], [("108", e["cust"], e["memo"], e["total"], 0), ("404", e["cust"], e["memo"], 0, e["supply"]),
                  ("255", e["cust"], e["memo"], 0, e["vat"])])
for x in ledger_coll:
    lines = [("103", x["cust"], x["memo"], x["ar"] - x["fee"], 0)]
    if x["fee"]:
        lines.append(("831", "", "송금수수료", x["fee"], 0))
    lines.append((x["acct"], x["cust"], x["memo"], 0, x["ar"]))
    J(x["date"], lines)

vat_q = defaultdict(lambda: [0, 0])  # 분기별 [매출세액, 매입세액]
for e in ledger_sales:
    vat_q[(e["date"].year, (e["date"].month - 1) // 3)][0] += e["vat"]

for y, m in months(2025, 1, 2026, 2):
    me = month_end(y, m)
    sales_m = sum(e["supply"] for e in ledger_sales if (e["date"].year, e["date"].month) == (y, m))
    # 원재료 매입 (월말 일괄, 익월말 지급)
    for s in SUPPLIERS:
        amt = int(sales_m * rng.uniform(0.07, 0.12) / 1000) * 1000
        vat = amt // 10
        J(me, [("153", "", f"원재료 매입 {s}", amt, 0), ("135", "", f"원재료 매입 {s}", vat, 0),
               ("251", "", f"원재료 매입 {s}", 0, amt + vat)])
        pd = next_biz(month_end(y + (m // 12), m % 12 + 1) - timedelta(days=3))
        if pd <= SUB_END:
            add_wd(pd, amt + vat, s)
            J(pd, [("251", "", f"외상대금 지급 {s}", amt + vat, 0), ("103", "", f"외상대금 지급 {s}", 0, amt + vat)])
        vat_q[(y, (m - 1) // 3)][1] += vat
    # 급여 25일
    gross = rng.randint(176, 186) * 1_000_000
    wh = int(gross * 0.095)
    pd = next_biz(date(y, m, 25))
    add_wd(pd, gross - wh, "급여이체", kind="대량이체")
    J(pd, [("801", "", f"{m}월 급여", gross, 0), ("254", "", f"{m}월 급여 원천징수", 0, wh),
           ("103", "", f"{m}월 급여", 0, gross - wh)])
    pd2 = next_biz(date(y + (m // 12), m % 12 + 1, 10))
    if pd2 <= SUB_END:
        add_wd(pd2, wh, "국세·지방세 납부", kind="CMS공동")
        J(pd2, [("254", "", f"{m}월 원천세 납부", wh, 0), ("103", "", f"{m}월 원천세 납부", 0, wh)])
    ins = rng.randint(16, 19) * 1_000_000
    add_wd(next_biz(date(y, m, 10)), ins, "국민건강보험공단", kind="CMS공동")
    J(next_biz(date(y, m, 10)), [("811", "", "4대보험료", ins, 0), ("103", "", "4대보험료", 0, ins)])
    fee = rng.randint(3, 8) * 100_000
    add_wd(next_biz(date(y, m, 20)), fee, "(주)가상세무회계", kind="인터넷")
    J(next_biz(date(y, m, 20)), [("831", "", "기장·자문 수수료", fee, 0), ("103", "", "기장·자문 수수료", 0, fee)])
    if m in (3, 6, 9, 12):
        intr = rng.randint(800, 1600) * 1000
        add_dep(me if me.weekday() < 5 else next_biz(me), intr, "결산이자", kind="이자")
        J(me, [("103", "", "예금이자", intr, 0), ("901", "", "예금이자", 0, intr)])

# 부가세 납부 (분기 다음달 25일)
for (y, q), (out_v, in_v) in sorted(vat_q.items()):
    pd = next_biz(date(y + (q == 3), (q * 3 + 4 - 1) % 12 + 1, 25))
    if pd <= SUB_END and out_v > in_v:
        add_wd(pd, out_v - in_v, "부가가치세 납부", kind="CMS공동")
        J(pd, [("255", "", "부가세 납부", out_v, 0), ("135", "", "부가세 납부", 0, in_v),
               ("103", "", "부가세 납부", 0, out_v - in_v)])

# 결산 분개 (2025-12-31)
purch = sum(l[3] for d, ls in journal if d.year == 2025 for l in ls if l[0] == "153")
J(FY_END, [("455", "", "제품매출원가 대체", int(purch * 0.97), 0), ("153", "", "제품매출원가 대체", 0, int(purch * 0.97))])
J(FY_END, [("818", "", "감가상각비 계상", 210_000_000, 0), ("207", "", "감가상각비 계상", 0, 210_000_000)])

# ---------------------------------------------------------------- 출력
COMPANY_DIR.mkdir(parents=True, exist_ok=True)


def sort_journal(year_from, year_to):
    rows = sorted([(d, ls) for d, ls in journal if year_from <= d <= year_to], key=lambda x: x[0])
    out, seq, last = [], 0, None
    for d, ls in rows:
        seq = seq + 1 if d == last else 1
        last = d
        for k, (a, cu, memo, dr, cr) in enumerate(ls):
            out.append((d, f"{seq:05d}", "차변" if dr else "대변", a, ACCT[a], cu, C[cu]["name"] if cu else "",
                        memo, dr, cr))
    return out


def write_journal(path, rows, period):
    wb = Workbook()
    ws = wb.active
    ws.title = "분개장"
    ws["A1"] = "분 개 장"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"회사명 : {COMPANY['name']}"
    ws["A3"] = f"기  간 : {period}"
    header_row(ws, 5, ["일자", "번호", "구분", "계정코드", "계정과목", "거래처코드", "거래처명", "적요", "차변", "대변"])
    for i, r in enumerate(rows, 6):
        body_row(ws, i, [d2s(r[0])] + list(r[1:]))
    widths(ws, [12, 8, 6, 9, 14, 10, 18, 34, 15, 15])
    ws.freeze_panes = "A6"
    wb.save(path)


j25 = sort_journal(FY_START, FY_END)
j26 = sort_journal(date(2026, 1, 1), SUB_END)
write_journal(COMPANY_DIR / "02_분개장_2025.xlsx", j25, "2025년 01월 01일 ~ 2025년 12월 31일")
write_journal(COMPANY_DIR / "03_분개장_2026_1-2월.xlsx", j26, "2026년 01월 01일 ~ 2026년 02월 28일")

# 시산표
OPEN_BAL = {"103": 1_250_000_000, "108": sum(opening.values()), "206": 2_400_000_000, "207": -780_000_000,
            "251": -420_000_000, "331": -1_000_000_000}
OPEN_BAL["375"] = -sum(OPEN_BAL.values())
mov = defaultdict(lambda: [0, 0])
for r in j25:
    mov[r[3]][0] += r[8]
    mov[r[3]][1] += r[9]
wb = Workbook()
ws = wb.active
ws.title = "합계잔액시산표"
ws["A1"] = "합 계 잔 액 시 산 표"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = f"회사명 : {COMPANY['name']}"
ws["A3"] = "기  간 : 2025년 12월 31일 현재"
header_row(ws, 5, ["차      변", "", "계 정 과 목", "대      변", ""])
ws.merge_cells("A5:B5"); ws.merge_cells("D5:E5"); ws.merge_cells("C5:C6")
for col, v in ((1, "잔  액"), (2, "합  계"), (4, "합  계"), (5, "잔  액")):
    ws.cell(6, col, v).font = BOLD
row = 7
tb_out = {}
for a in sorted(ACCT):
    ob = OPEN_BAL.get(a, 0)
    dr = mov[a][0] + max(ob, 0)
    cr = mov[a][1] + max(-ob, 0)
    if dr == 0 and cr == 0:
        continue
    bal = dr - cr
    tb_out[a] = bal
    body_row(ws, row, [max(bal, 0), dr, f"{ACCT[a]}", cr, max(-bal, 0)])
    row += 1
tdr = sum(ws.cell(r, 2).value for r in range(7, row))
tcr = sum(ws.cell(r, 4).value for r in range(7, row))
body_row(ws, row, [sum(ws.cell(r, 1).value for r in range(7, row)), tdr, "합      계", tcr,
                   sum(ws.cell(r, 5).value for r in range(7, row))])
for c_ in range(1, 6):
    ws.cell(row, c_).font = BOLD
widths(ws, [18, 18, 18, 18, 18])
wb.save(COMPANY_DIR / "01_합계잔액시산표_2025.xlsx")
assert tdr == tcr, (tdr, tcr)

# 계정별원장: 외상매출금 (월계·누계 행 포함)
wb = Workbook()
ws = wb.active
ws.title = "108 외상매출금"
ws["A1"] = "계 정 별 원 장"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = f"회사명 : {COMPANY['name']}"
ws["A3"] = "기  간 : 2025년 01월 01일 ~ 2025년 12월 31일        계정과목 : 108 외상매출금"
header_row(ws, 5, ["일자", "적요", "코드", "거래처", "차변", "대변", "잔액"])
bal = OPEN_BAL["108"]
r = 6
body_row(ws, r, ["", "[전기이월]", "", "", OPEN_BAL["108"], 0, bal], fill=BAND_FILL)
r += 1
cum_dr, cum_cr = OPEN_BAL["108"], 0
for m in range(1, 13):
    mdr = mcr = 0
    for x in j25:
        if x[3] != "108" or int(x[0].month) != m:
            continue
        bal += x[8] - x[9]
        mdr += x[8]; mcr += x[9]
        body_row(ws, r, [f"{x[0]:%m-%d}", x[7], x[5], x[6], x[8], x[9], bal])
        r += 1
    cum_dr += mdr; cum_cr += mcr
    body_row(ws, r, ["", "[ 월      계 ]", "", "", mdr, mcr, ""], fill=BAND_FILL); r += 1
    body_row(ws, r, ["", "[ 누      계 ]", "", "", cum_dr, cum_cr, ""], fill=BAND_FILL); r += 1
widths(ws, [8, 34, 8, 18, 15, 15, 16])
ws.freeze_panes = "A6"
wb.save(COMPANY_DIR / "04_계정별원장_외상매출금_2025.xlsx")

# 거래처원장 (외상매출금 거래처별 잔액)
wb = Workbook()
ws = wb.active
ws.title = "거래처원장"
ws["A1"] = "거 래 처 원 장"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = f"회사명 : {COMPANY['name']}"
ws["A3"] = "기  간 : 2025년 01월 01일 ~ 2025년 12월 31일        계정과목 : 108 외상매출금        [잔액]"
header_row(ws, 5, ["코드", "거래처", "등록번호", "대표자명", "전기이월", "차변", "대변", "잔액"])
per = defaultdict(lambda: [0, 0])
for x in j25:
    if x[3] == "108":
        per[x[5]][0] += x[8]; per[x[5]][1] += x[9]
r = 6
for c in sorted(customers, key=lambda c: c["code"]):
    ob = opening.get(c["code"], 0)
    dr, cr = per[c["code"]]
    if not (ob or dr or cr):
        continue
    body_row(ws, r, [c["code"], c["name"], c["bizno"], c["ceo"], ob, dr, cr, ob + dr - cr])
    r += 1
body_row(ws, r, ["", "합      계", "", "", f"=SUM(E6:E{r-1})", f"=SUM(F6:F{r-1})", f"=SUM(G6:G{r-1})", f"=SUM(H6:H{r-1})"],
         fill=BAND_FILL, num_cols=(5, 6, 7, 8))
widths(ws, [8, 18, 14, 10, 15, 15, 15, 15])
wb.save(COMPANY_DIR / "05_거래처원장_외상매출금_2025.xlsx")

# 거래처등록
wb = Workbook()
ws = wb.active
ws.title = "거래처등록"
ws["A1"] = "거 래 처 등 록 현 황 (일반거래처)"
ws["A1"].font = Font(bold=True, size=14)
header_row(ws, 3, ["코드", "거래처명", "사업자등록번호", "대표자", "업태", "종목", "주소", "담당자 이메일"])
for i, c in enumerate(customers, 4):
    body_row(ws, i, [c["code"], c["name"], c["bizno"], c["ceo"], "제조", "기계부품", c["addr"], c["email"]])
widths(ws, [8, 18, 15, 10, 8, 10, 34, 28])
wb.save(COMPANY_DIR / "06_거래처등록현황.xlsx")

# 홈택스 매출 전자세금계산서 목록 (실측 기록 기준 33열, 1행 회사·3행 총계·5행 띠·6행 머리글)
HT_HEAD = ["작성일자", "승인번호", "발급일자", "전송일자", "공급자사업자등록번호", "종사업장번호", "상호", "대표자명", "주소",
           "공급받는자사업자등록번호", "종사업장번호", "상호", "대표자명", "주소", "합계금액", "공급가액", "세액",
           "전자세금계산서분류", "전자세금계산서종류", "발급유형", "비고", "영수/청구 구분", "공급자 이메일",
           "공급받는자 이메일1", "공급받는자 이메일2", "품목일자", "품목명", "품목규격", "품목수량", "품목단가",
           "품목공급가액", "품목세액", "품목비고"]


def write_hometax(path, rows, label):
    wb = Workbook()
    ws = wb.active
    ws.title = "세금계산서"
    ws.append(["사업자 등록번호", COMPANY["bizno"], "상호", COMPANY["name"], "대표자 명", COMPANY["ceo"]])
    ws.append([])
    ws.append(["총 합계금액", sum(i["total"] for i in rows), "총 공급가액", sum(i["supply"] for i in rows),
               "총 세액", sum(i["vat"] for i in rows)])
    ws.append([])
    band = ["매출 전자(수정)세금계산서 목록조회 " + label] + [""] * 24 + ["대표 품목"]
    ws.append(band)
    header_row(ws, 6, HT_HEAD)
    items_ws = wb.create_sheet("품목")
    items_ws.append(["매출 전자(수정)세금계산서 품목 " + label]); items_ws.append([]); items_ws.append([])
    header_row(items_ws, 4, ["승인번호", "품목순번", "공급자등록번호", "공급받는자등록번호", "일자", "품목명", "규격",
                             "수량", "단가", "공급가액", "세액", "비고"])
    for inv in sorted(rows, key=lambda i: (i["wdate"], i["appr"])):
        c = C[inv["cust"]]
        it = inv["items"][0]
        ws.append([d2s(inv["wdate"]), inv["appr"], d2s(inv["idate"]), d2s(inv["sdate"]), COMPANY["bizno"], "",
                   COMPANY["name"], COMPANY["ceo"], COMPANY["addr"], c["bizno"], "", c["name"], c["ceo"], c["addr"],
                   inv["total"], inv["supply"], inv["vat"],
                   "수정세금계산서" if inv["kind"] == "수정" else "세금계산서",
                   "환입" if inv["kind"] == "수정" else "일반", "정발급",
                   inv["note"], "청구", COMPANY["email"], c["email"], "",
                   d2s(it.get("date", inv["wdate"])), it["name"] + (f" 외 {len(inv['items']) - 1}건" if len(inv["items"]) > 1 else ""),
                   it["spec"], it["qty"], it["price"], it["supply"], int(it["supply"] * 0.1) if it["supply"] > 0 else -int(-it["supply"] * 0.1), ""])
        for k, it in enumerate(inv["items"], 1):
            items_ws.append([inv["appr"], k, COMPANY["bizno"], c["bizno"], d2s(it.get("date", inv["wdate"])),
                             it["name"], it["spec"], it["qty"], it["price"], it["supply"],
                             int(it["supply"] * 0.1) if it["supply"] > 0 else -int(-it["supply"] * 0.1), ""])
    widths(ws, [11, 30, 11, 11, 14, 6, 14, 8, 20, 14, 6, 16, 8, 20, 13, 13, 12, 14, 10, 8, 40, 8, 22, 22, 8, 11, 22, 8, 8, 9, 13, 11, 8])
    ws.freeze_panes = "A7"
    wb.save(path)


inv25 = [i for i in invoices if i["wdate"].year == 2025]
inv26 = [i for i in invoices if i["wdate"].year == 2026]
write_hometax(COMPANY_DIR / "07_매출전자세금계산서목록_2025.xlsx", inv25, "(작성일자 2025-01-01 ~ 2025-12-31)")
write_hometax(COMPANY_DIR / "08_매출전자세금계산서목록_2026_1-2월.xlsx", inv26, "(작성일자 2026-01-01 ~ 2026-02-28)")

# 은행 거래내역 (KB국민은행 형식: 거래일시|적요|기재내용|찾으신금액|맡기신금액|잔액|거래점)
txs = [(d["dt"], d["kind"], d["payer"][:10], 0, d["amount"]) for d in deposits] + \
      [(w["dt"], w["kind"], w["payee"][:10], w["amount"], 0) for w in withdrawals]
txs.sort(key=lambda t: t[0])
wb = Workbook()
ws = wb.active
ws.title = "거래내역조회"
ws["A1"] = "거래내역조회"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = "계좌번호 : 123456-01-987654 (기업자유예금)"
ws["A3"] = f"예금주명 : {COMPANY['name']}"
ws["A4"] = "조회기간 : 2025.01.01 ~ 2026.02.28"
header_row(ws, 6, ["거래일시", "적요", "기재내용", "찾으신금액", "맡기신금액", "잔액", "거래점"])
bal = OPEN_BAL["103"]
for i, (dt, kind, memo, out_, in_) in enumerate(txs, 7):
    bal += in_ - out_
    body_row(ws, i, [dt.strftime("%Y.%m.%d %H:%M:%S"), kind, memo, out_, in_, bal,
                     "인터넷" if kind in ("인터넷", "타행이체") else "안산중앙"])
widths(ws, [20, 10, 16, 15, 15, 16, 10])
ws.freeze_panes = "A7"
wb.save(COMPANY_DIR / "09_은행거래내역_KB국민_보통예금.xlsx")

# 부가세 신고서 과세표준 요약 (회사가 홈택스 발급분 기준으로 신고)
wb = Workbook()
ws = wb.active
ws.title = "과세표준"
ws["A1"] = "부가가치세 신고서 과세표준 및 매출세액 요약 (2025년, 회사 신고분 전기)"
ws["A1"].font = Font(bold=True, size=13)
header_row(ws, 3, ["신고구분", "과세기간", "과세 세금계산서 발급분 금액", "세액", "과세 기타 금액", "영세율 금액", "합계 금액"])
periods = [("1기 예정", (1, 3)), ("1기 확정", (4, 6)), ("2기 예정", (7, 9)), ("2기 확정", (10, 12))]
for k, (lab, (m0, m1)) in enumerate(periods, 4):
    s = sum(i["supply"] for i in inv25 if m0 <= i["wdate"].month <= m1)
    v = sum(i["vat"] for i in inv25 if m0 <= i["wdate"].month <= m1)
    body_row(ws, k, [lab, f"2025.{m0:02d}.01~2025.{m1:02d}.{month_end(2025, m1).day}", s, v, 0, 0, s])
widths(ws, [10, 24, 22, 16, 14, 12, 18])
wb.save(COMPANY_DIR / "10_부가가치세신고서_과세표준요약_2025.xlsx")

# 정답지와 내부 메타 (엔진은 읽지 않음)
customer_books = {}  # 거래처가 자기 장부에 기록한 채무 잔액 (채권채무조회 회신 시뮬레이션용)
for c in customers:
    if c is ghost:
        continue
    billed = opening.get(c["code"], 0) + sum(i["total"] for i in inv25 if i["cust"] == c["code"])
    if c is c2:
        billed += 8_030_000
    paid = sum(d["amount"] + (1_000 if c["pay"] == "fee" and d["links"] and not isinstance(d["links"][0], tuple) else 0)
               for d in deposits if d["dt"].date() <= FY_END and d["payer"] == c["payer"])
    if c is c9:
        paid += sum(d["amount"] for d in deposits if d["payer"] == "태성홀딩스" and d["dt"].date() <= FY_END)
    customer_books[c["code"]] = billed - paid

(DATA / "truth.json").write_text(json.dumps({
    "note": "심어 둔 이상 거래 정답지. K-Tieout 엔진은 이 파일을 읽지 않는다. 비교(compare.py)에만 쓴다.",
    "anomalies": truth,
    "customer_books": customer_books,
    "ghost_customer": ghost["code"],
    "tb": tb_out,
    "counts": {"invoices_2025": len(inv25), "ledger_sales_2025": sum(1 for e in ledger_sales if e["date"].year == 2025),
               "deposits": len(deposits), "withdrawals": len(withdrawals)},
}, ensure_ascii=False, indent=1, default=str))

print("세금계산서 2025:", len(inv25), "건 / 2026 1-2월:", len(inv26), "건")
print("원장 매출 2025:", sum(1 for e in ledger_sales if e["date"].year == 2025), "건,",
      f"{sum(e['supply'] for e in ledger_sales if e['date'].year == 2025):,}원")
print("은행 입금", len(deposits), "건 / 출금", len(withdrawals), "건")
print("시산표 외상매출금", f"{tb_out['108']:,}", "제품매출", f"{-tb_out['404']:,}", "보통예금", f"{tb_out['103']:,}")
