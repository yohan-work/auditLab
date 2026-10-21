"""As-Is: 지금 방식(표본 기반)으로 1~3년차 회계사가 작성하는 매출·매출채권 조서를 재현한다.

감사마스터(중소회계법인협의회) 조서 양식의 항목을 따른다.
  4000C-1 총괄표및Recap-매출채권   4000C-2 계정내역,명세대사   5050 채권채무조회 표본샘플링
  5051 일반채권채무조회서          4000C-3 조회서대사          4000P-2 부가세대사표
  4000P-5 매출Cut-off              4000P-10 거래내역-증빙 검토

회사 제출 자료로 볼 수 있는 것만 쓴다. 단, 외부에서 받는 증거 두 가지는 정답지에서 가져온다:
  - 채권채무조회 회신 금액 (거래처가 자기 장부로 회신) ← truth.json customer_books
  - 출고증의 출고일 ← 홈택스 품목 시트의 품목일자로 대신함
"""
import json
import random
from collections import Counter
from datetime import date

from openpyxl import Workbook, load_workbook

from common import (BAND_FILL, BOLD, COMPANY_DIR, DATA, FLAG_FILL, FY_END, MATERIALITY, OUT, PERFORMANCE_MAT,
                    TICKMARKS, TRIVIAL, WARN_FILL, body_row, d2s, header_row, widths, wp_header)
from ktieout import P, read_customers, read_hometax, read_journal, read_opening, read_vat_return, read_bank

PREP, REV = "회계사 A (2년차)", "매니저 B"
PD, RD = "2026.02.16", "2026.02.20"
SAMPLE_N = 25
CUTOFF_MIN = 10_000_000
CONF_COVER = 0.70


def tick_legend(ws, r):
    ws.cell(r, 1, "Tick Mark").font = BOLD
    for k, (t, desc) in enumerate(TICKMARKS + [("V", "증빙(세금계산서·출고증·입금)과 대사함"), ("X", "예외사항 (아래 메모 참조)")], r + 1):
        ws.cell(k, 1, t)
        ws.cell(k, 2, desc)
    return k + 2


def main():
    rng = random.Random(7)
    truth = json.load(open(DATA / "truth.json"))
    custs = read_customers()
    j = read_journal("02_분개장_2025.xlsx")
    j26 = read_journal("03_분개장_2026_1-2월.xlsx")
    inv = read_hometax("07_매출전자세금계산서목록_2025.xlsx") + read_hometax("08_매출전자세금계산서목록_2026_1-2월.xlsx")
    opening, ledger_bal = read_opening()
    vat_ret = read_vat_return()
    bank = read_bank()

    item_dates = {}
    for f in ("07_매출전자세금계산서목록_2025.xlsx", "08_매출전자세금계산서목록_2026_1-2월.xlsx"):
        ws = load_workbook(COMPANY_DIR / f)["품목"]
        for r in ws.iter_rows(min_row=5, values_only=True):
            if r[0]:
                item_dates.setdefault(r[0], P(r[4]))

    def sales_of(rows):
        out = []
        for r in rows:
            if r["acct"] == "108" and r["dr"]:
                out.append({"date": r["date"], "no": r["no"], "cust": r["cust"], "total": r["dr"], "memo": r["memo"]})
        return out

    s25, s26 = sales_of(j), sales_of(j26)
    out = OUT / "as-is"
    out.mkdir(parents=True, exist_ok=True)
    findings, effort = [], Counter()

    def F(proc, cust, level, desc, amount=0):
        findings.append({"proc": proc, "cust": cust, "cust_name": custs[cust]["name"] if cust else "", "level": level,
                         "desc": desc, "amount": amount})

    def find_inv(s):
        """회계사가 세금계산서철(홈택스 목록)에서 같은 거래처·같은 금액 건을 찾는다."""
        c = custs[s["cust"]]["bizno"]
        cands = [i for i in inv if i["bizno"] == c and i["total"] == s["total"] and abs((i["wdate"] - s["date"]).days) <= 31]
        if cands:
            return min(cands, key=lambda i: abs((i["wdate"] - s["date"]).days)), "일치"
        near = [i for i in inv if i["bizno"] == c and abs((i["wdate"] - s["date"]).days) <= 5 and i["total"] > 0]
        if near:
            return min(near, key=lambda i: abs(i["total"] - s["total"])), "금액 불일치"
        return None, "없음"

    def find_dep(cust, amount, after):
        n = custs[cust]["name"].replace("(주)", "").replace("주식회사", "")
        return [b for b in bank if b["in"] and n[:3] in (b["memo"] or "") and b["dt"].date() >= after
                and abs(b["in"] - amount) <= 3_000]

    # ---------------- 4000C-1 총괄표
    wb = Workbook()
    ws = wb.active
    ws.title = "4000C-1"
    r = wp_header(ws, "4000C-1", "총괄표및Recap-매출채권(장단기)", PREP, REV, PD, RD)
    header_row(ws, r, ["계정과목", "전기말(감사후)", "당기말 회사제시", "수정분개", "당기말 감사후", "증감", "Ref.", "Tick-Mark"])
    ar = sum(ledger_bal.values())
    body_row(ws, r + 1, ["외상매출금", sum(opening.values()), ar, 0, ar, ar - sum(opening.values()), "4000C-2", "PY, PTB"])
    body_row(ws, r + 2, ["대손충당금", 0, 0, 0, 0, 0, "4000C-4", "PY, PTB"])
    ws.cell(r + 4, 1, f"중요성: 전체 {MATERIALITY:,} / 수행 {PERFORMANCE_MAT:,} / 명백히 사소한 금액 {TRIVIAL:,} (2700A 참조)")
    widths(ws, [16, 16, 16, 12, 16, 16, 9, 12])
    effort["총괄표 작성 (시산표·전기조서 대사)"] += 1

    # ---------------- 4000C-2 명세대사
    ws = wb.create_sheet("4000C-2")
    r = wp_header(ws, "4000C-2", "계정내역,명세대사", PREP, REV, PD, RD,
                  "회사제시 거래처원장 잔액명세의 합계를 시산표와 대사 (거래처별 잔액의 실재성은 5050·5051 조회로 확인)")
    header_row(ws, r, ["계정", "거래처", "기초금액", "증가", "감소", "기말잔액", "Tick-Mark", "검토 No", "비고"])
    wsl = load_workbook(COMPANY_DIR / "05_거래처원장_외상매출금_2025.xlsx").active
    k = r
    for row in wsl.iter_rows(min_row=6, values_only=True):
        if not row[0]:
            continue
        k += 1
        body_row(ws, k, ["외상매출금", row[1], row[4], row[5], row[6], row[7], "", "", ""])
        effort["명세 거래처 행 검토"] += 1
    body_row(ws, k + 1, ["", "합      계", sum(opening.values()), "", "", ar, "F, PTB", "", "시산표와 일치"], fill=BAND_FILL)
    tick_legend(ws, k + 3)
    widths(ws, [11, 18, 15, 15, 15, 15, 10, 8, 18])

    # ---------------- 5050 채권채무조회 표본샘플링 (금액 기준 상위 → 커버리지 70%)
    ws = wb.create_sheet("5050")
    r = wp_header(ws, "5050", "채권채무조회 표본샘플링", PREP, REV, PD, RD,
                  f"계정 선택: 외상매출금 / 잔액 있는 것만 / 금액 기준 상위 거래처부터 금액추출율 {CONF_COVER:.0%} 이상")
    bal_sorted = sorted([(c, b) for c, b in ledger_bal.items() if b > 0], key=lambda x: -x[1])
    sel, acc = [], 0
    for c, b in bal_sorted:
        if acc < CONF_COVER * ar:
            sel.append(c)
            acc += b
    header_row(ws, r, ["계정", "거래처", "기초금액", "차변금액", "대변금액", "기말잔액", "통화", "외화잔액", "선택"])
    for k, row in enumerate(wsl.iter_rows(min_row=6, values_only=True), r + 1):
        if not row[0]:
            continue
        body_row(ws, k, ["외상매출금", row[1], row[4], row[5], row[6], row[7], "KRW", 0, "V" if row[0] in sel else ""])
    k = ws.max_row + 2
    pos = [b for _, b in bal_sorted]
    ws.cell(k, 1, f"· 계정금액/ 추출금액/ 금액추출율 : {ar:,} / {acc:,} / {acc / ar:.1%}")
    ws.cell(k + 1, 1, f"· 전체표본수/ 추출표본수/ 표본수추출율 : {len(pos)} / {len(sel)} / {len(sel) / len(pos):.1%}")
    widths(ws, [11, 18, 15, 15, 15, 15, 6, 8, 6])
    effort["조회 대상 선정"] += 1

    # ---------------- 5051 조회서 발송·회신 / 4000C-3 조회서대사
    books = truth["customer_books"]
    ws = wb.create_sheet("5051")
    r = wp_header(ws, "5051", "일반채권채무조회서", PREP, REV, PD, RD, "KICPA·NICE 채권채무 전자조회로 발송, 회신은 감사인이 직접 수령")
    header_row(ws, r, ["발신여부", "회신여부", "거래처", "계정", "기말잔액", "외화잔액 (통화)", "조회서 No", "조회금액",
                       "외화 조회금액", "차이", "외화차이"])
    conf = []
    for k, c in enumerate(sel, 1):
        reply = books.get(c)
        bal = ledger_bal[c]
        diff = None if reply is None else bal - reply
        conf.append((c, bal, reply, diff, f"AR-{k:03d}"))
        body_row(ws, r + k, ["발신", "회신" if reply is not None else "미회신", custs[c]["name"], "외상매출금", bal, "",
                             f"AR-{k:03d}", reply if reply is not None else "", "", diff if diff is not None else "", ""],
                 fill=None if diff == 0 else WARN_FILL)
        effort["조회서 발송·회신 관리"] += 1
    widths(ws, [8, 8, 18, 11, 15, 10, 9, 15, 10, 14, 8])

    ws = wb.create_sheet("4000C-3")
    r = wp_header(ws, "4000C-3", "조회서대사-매출채권(장단기)", PREP, REV, PD, RD)
    ws.cell(r, 1, f"01. Test 대상 모집단  (1) 모집단 총금액 : {ar:,}  (2) 회사제시 자료명 거래처원장(외상매출금)")
    ws.cell(r + 1, 1, "(3) 모집단의 완전성 확인: 제시 명세서가 재무제표/시산표 및 총괄표의 금액과 일치함을 확인함")
    r += 3
    replied = [x for x in conf if x[2] is not None]
    header_row(ws, r, ["계정", "구분", "전체표본", "조회대상 Sampling", "Sample/전체 비율", "조회 회신", "Sample중 회신율",
                       "전체대비 회신율", "비고"])
    body_row(ws, r + 1, ["외상매출금", "금액", ar, acc, f"{acc / ar:.1%}", sum(x[1] for x in replied),
                         f"{sum(x[1] for x in replied) / acc:.1%}", f"{sum(x[1] for x in replied) / ar:.1%}", ""])
    body_row(ws, r + 2, ["외상매출금", "건수", len(pos), len(sel), f"{len(sel) / len(pos):.1%}", len(replied),
                         f"{len(replied) / len(sel):.1%}", f"{len(replied) / len(pos):.1%}", ""])
    r += 4
    header_row(ws, r, ["NO", "거래처", "금액", "통화", "외화금액", "회신여부", "조회서번호", "조회확인 금액", "조회확인 외화", "차이",
                       "외화차이", "Tick-Mark", "검토 No"])
    for k, (c, bal, reply, diff, no) in enumerate(conf, 1):
        body_row(ws, r + k, [k, custs[c]["name"], bal, "KRW", "", "회신" if reply is not None else "미회신", no,
                             reply if reply is not None else "", "", diff if diff is not None else "", "",
                             "V" if diff == 0 else "X", "" if diff == 0 else f"C3-{k:02d}"],
                 fill=None if diff == 0 else FLAG_FILL)
        effort["회신 금액 대사"] += 1
        if reply is None:
            F("4000C-3", c, "징후", f"미회신 → 대체적 절차(후속 입금·증빙 확인) 필요", bal)
            effort["미회신 대체적 절차 (후속입금·증빙 수작업 확인)"] += 1
        elif abs(diff) > 0:
            F("4000C-3", c, "징후", f"조회 차이 {diff:,}원 → 회사에 차이 내역 소명 요청 (원인 미규명)", diff)
            effort["조회 차이 원인 조사 (회사 소명 요청·대조)"] += 1
    k = r + len(conf) + 2
    ws.cell(k, 1, "기타검토사항 / 결론").font = BOLD
    ws.cell(k + 1, 1, "차이 발생 거래처는 회사에 차이조정표를 요청함 (2026.02.18 요청, 회신 대기). 미회신 거래처는 대체적 절차 수행 예정.")
    widths(ws, [5, 18, 15, 6, 8, 8, 10, 15, 8, 14, 8, 9, 8])

    # ---------------- 4000P-2 부가세대사표 (분기 총액)
    ws = wb.create_sheet("4000P-2")
    r = wp_header(ws, "4000P-2", "부가세대사표-매출및기타수익", PREP, REV, PD, RD,
                  "· 테스트 절차/방법: 부가세 신고서 과세표준과 손익계산서 매출액을 분기별로 대사")
    header_row(ws, r, ["공급가액", "1/4분기", "2/4분기", "3/4분기", "4/4분기", "계", "검토사항"])
    q = lambda d: (d.month - 1) // 3
    sup = {}
    for x in j:
        if x["acct"] == "404":
            sup[q(x["date"])] = sup.get(q(x["date"]), 0) + x["cr"]
    gl_q = [sup.get(i, 0) for i in range(4)]
    rows = [("(매출 부가세 신고) 과세(계산서)", vat_ret), ("영세율", [0] * 4), ("기타과세", [0] * 4),
            ("과세 합계", vat_ret), ("면세(계산서)", [0] * 4), ("면세 합계", [0] * 4), ("전체 합계", vat_ret),
            ("가산(차감) : 매출조정항목", [0] * 4), ("  거래시기차이(선수금 수령)", [0] * 4), ("  유형자산 처분액", [0] * 4),
            ("가산(차감)후 금액", vat_ret), ("손익계산서 매출액", gl_q),
            ("차이", [g - v for g, v in zip(gl_q, vat_ret)])]
    for k, (lab, vals) in enumerate(rows, r + 1):
        body_row(ws, k, [lab] + list(vals) + [sum(vals), "회사 소명 요청" if lab == "차이" and any(vals) else ""],
                 fill=FLAG_FILL if lab == "차이" and any(vals) else None)
    tot = sum(g - v for g, v in zip(gl_q, vat_ret))
    ws.cell(r + len(rows) + 2, 1, f"· 테스트 결과 메모, 결론: 연간 차이 {tot:,}원 (수행중요성 {PERFORMANCE_MAT:,} 대비 "
            f"{abs(tot) / PERFORMANCE_MAT:.0%}). 회사에 매출조정항목 명세 요청함. 원인 규명 전 결론 보류.")
    widths(ws, [32, 16, 16, 16, 16, 18, 14])
    effort["부가세대사 (분기 합계 작성)"] += 4
    F("4000P-2", None, "총액", f"신고서 vs 손익계산서 연간 차이 {tot:,}원 인지 → 원인 분해는 회사 소명에 의존", tot)
    effort["부가세 차이 원인 조사 (회사 소명 요청·대조)"] += 1

    # ---------------- 4000P-5 매출 Cut-off (금기 12/24~31, 차기 1/1~1/9, 1천만원 이상)
    ws = wb.create_sheet("4000P-5")
    r = wp_header(ws, "4000P-5", "매출Cut-off-매출및기타수익", PREP, REV, PD, RD,
                  f"· 테스트 절차/방법: 결산일 전후 5영업일 매출 중 {CUTOFF_MIN:,}원 이상 전 건, 원시 출고증빙의 출고일과 회계기간 비교")
    cut_cur = [s for s in s25 if s["date"] >= date(2025, 12, 24) and s["total"] >= CUTOFF_MIN]
    cut_nxt = [s for s in s26 if s["date"] <= date(2026, 1, 9) and s["total"] >= CUTOFF_MIN]
    k = r
    for title, lst, yr in (("1. 금기매출", cut_cur, 2025), ("2. 차기매출", cut_nxt, 2026)):
        ws.cell(k, 1, title).font = BOLD
        header_row(ws, k + 1, ["No.", "전표일", "전표번호", "내역", "거래처", "품목", "수량", "금액", "첨부증빙",
                               "증빙일(원시출고증빙)", "회계기간 일치여부"])
        k += 2
        for n, s in enumerate(lst, 1):
            i, _ = find_inv(s)
            ship = item_dates.get(i["appr"]) if i else None
            ok = ship is not None and ship.year == yr
            body_row(ws, k, [n, d2s(s["date"]), s["no"], s["memo"], custs[s["cust"]]["name"], s["memo"].split(" ")[1] if " " in s["memo"] else "",
                             "", s["total"], "출고증·세금계산서" if i else "증빙 미제시", d2s(ship) if ship else "",
                             "일치" if ok else "불일치"], fill=None if ok else FLAG_FILL)
            effort["Cut-off 증빙 대조 (출고증·세금계산서)"] += 1
            if not ok:
                F("4000P-5", s["cust"], "발견", f"{title[3:]} {d2s(s['date'])} 전표, 출고일 {d2s(ship) if ship else '증빙 없음'}"
                  " → 회계기간 불일치", s["total"])
            k += 1
        k += 1
    ws.cell(k, 1, "· 테스트 결과 메모, 결론: 위 불일치 건은 회사에 수정 여부 확인 요청")
    widths(ws, [5, 11, 8, 28, 16, 12, 6, 14, 16, 14, 12])

    # ---------------- 4000P-10 거래내역-증빙 검토 (MUS 25건)
    ws = wb.create_sheet("4000P-10")
    tot_amt = sum(s["total"] for s in s25)
    interval = tot_amt / SAMPLE_N
    start = rng.uniform(0, interval)
    points = [start + interval * n for n in range(SAMPLE_N)]
    cum, pick, pi = 0, [], 0
    for s in s25:
        lo, cum = cum, cum + s["total"]
        while pi < len(points) and lo <= points[pi] < cum:
            if not pick or pick[-1] is not s:
                pick.append(s)
            pi += 1
    r = wp_header(ws, "4000P-10", "거래내역-증빙 검토-매출및기타수익", PREP, REV, PD, RD,
                  f"· 테스트 절차/방법: 2025년 매출 전표 모집단 {len(s25)}건·{tot_amt:,}원에서 금액단위표본추출(MUS) "
                  f"{SAMPLE_N}건 (표본간격 {interval:,.0f}원). 세금계산서·출고증·입금 증빙과 대사")
    header_row(ws, r, ["No.", "전표일", "전표번호", "거래처", "금액", "세금계산서 승인번호", "작성일", "세금계산서 금액",
                       "출고일", "입금일", "입금액", "일치여부", "Tick-Mark", "비고"])
    for n, s in enumerate(pick, 1):
        i, how = find_inv(s)
        ship = item_dates.get(i["appr"]) if i else None
        dep = find_dep(s["cust"], i["total"] if i else s["total"], s["date"])[:1]
        ok = how == "일치" and dep
        note = "" if ok else ("세금계산서 없음" if how == "없음" else ("금액 불일치" if how == "금액 불일치" else "입금 미확인(후속기간)"))
        body_row(ws, r + n, [n, d2s(s["date"]), s["no"], custs[s["cust"]]["name"], s["total"], i and i["appr"],
                             i and d2s(i["wdate"]), i and i["total"], d2s(ship) if ship else "",
                             f"{dep[0]['dt']:%Y-%m-%d}" if dep else "", dep[0]["in"] if dep else "",
                             "일치" if ok else "예외", "V" if ok else "X", note], fill=None if ok else FLAG_FILL)
        effort["증빙 대조 (세금계산서·출고증·통장 각 1건)"] += 3
        if not ok:
            F("4000P-10", s["cust"], "발견" if how != "일치" else "징후", f"표본 {n}: {note}", s["total"])
    k = r + len(pick) + 2
    covered = sum(s["total"] for s in pick)
    ws.cell(k, 1, f"· 테스트 결과: 표본 {len(pick)}건 / {covered:,}원 (모집단 금액의 {covered / tot_amt:.1%}, 건수의 {len(pick) / len(s25):.1%})")
    tick_legend(ws, k + 2)
    widths(ws, [5, 11, 8, 16, 14, 30, 11, 14, 11, 11, 14, 8, 6, 20])

    path = out / "As-Is_조서_샘플제조_2025.xlsx"
    wb.save(path)

    tested_ledger = {(d2s(s["date"]), s["cust"], s["total"]) for s in pick + cut_cur}
    json.dump({
        "findings": findings,
        "effort": dict(effort),
        "coverage": {"vouch_n": len(pick), "vouch_amt": covered, "pop_n": len(s25), "pop_amt": tot_amt,
                     "cutoff_n": len(cut_cur) + len(cut_nxt), "conf_n": len(sel), "conf_amt": acc, "ar": ar,
                     "conf_reply_n": len(replied), "conf_diff_n": sum(1 for x in conf if x[3])},
        "sample": [{"date": d2s(s["date"]), "cust": s["cust"], "total": s["total"]} for s in pick],
        "confirm": [{"cust": c, "ledger": b, "reply": rp, "diff": df} for c, b, rp, df, _ in conf],
        "vat_diff": tot,
    }, open(out / "asis_result.json", "w"), ensure_ascii=False, indent=1, default=str)
    return path, findings, effort


if __name__ == "__main__":
    path, findings, effort = main()
    print("조서:", path)
    for f in findings:
        print(f"  [{f['proc']}] {f['level']} {f['cust_name']}: {f['desc']}")
    print("작업량:", dict(effort))
