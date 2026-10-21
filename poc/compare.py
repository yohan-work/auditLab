"""As-Is(표본) vs K-Tieout(전수) 비교. 정답지(truth.json)로 채점한다.

결과: output/비교_As-Is_vs_K-Tieout.xlsx, output/metrics.json (데모 페이지용)
"""
import json
from datetime import date

from openpyxl import Workbook, load_workbook

from common import (BOLD, COMPANY_DIR, DATA, FLAG_FILL, OK_FILL, OUT, WARN_FILL, body_row, header_row, widths)

truth = json.load(open(DATA / "truth.json"))
asis = json.load(open(OUT / "as-is" / "asis_result.json"))
tobe = json.load(open(OUT / "to-be" / "ktieout_result.json"))

KT_KIND = {"A1": "원장O·세금계산서X·입금X", "A2": "원장O·세금계산서X·입금O", "A3": "세금계산서O·원장X",
           "A4": "금액 불일치", "A5": "수정세금계산서 원장 미반영", "A6": "기간귀속 (당기 과대)",
           "A7": "기간귀속 (당기 과소)", "A8": "기말 직전 매출 후속 취소", "A9": "제3자 명의 입금",
           "A10": "원장 중복 계상 의심", "A11": "장기 미회수"}

sample = {(s["date"], s["cust"], s["total"]) for s in asis["sample"]}
sample_cust = {s["cust"] for s in asis["sample"]}
conf = {c["cust"]: c for c in asis["confirm"]}
cut = {(f["cust"]) for f in asis["findings"] if f["proc"] == "4000P-5"}


def asis_result(a):
    """지금 방식으로 이 이상 거래가 드러나는가. (결과, 근거 조서, 설명)"""
    i, c = a["id"], a["cust"]
    cf = conf.get(c)
    diff = cf and cf["reply"] is not None and cf["diff"]
    if i in ("A1", "A2"):
        if (a["ledger_date"], c, a["amount"]) in sample:
            return "발견", "4000P-10", "증빙 대조 표본에 포함되어 세금계산서 없음 확인"
        if cf and cf["reply"] is None:
            return "징후", "4000C-3", "조회 미회신 → 대체적 절차에서 드러날 수 있음"
        return "미발견", "-", f"표본 {len(sample)}건에 포함되지 않음" + (", 조회 대상도 아님" if not cf else "")
    if i == "A3":
        return "미발견", "4000P-2", "원장에서 표본을 뽑는 방향이라 원장에 없는 매출은 표본에 안 잡힘. 부가세대사 총액 차이에 섞여 있음"
    if i == "A4":
        if c in sample_cust and any(s["cust"] == c and abs(s["total"] - 27_830_000) < 1 for s in asis["sample"]):
            return "발견", "4000P-10", "표본 대조에서 금액 불일치 확인"
        if diff:
            return "징후", "4000C-3", f"조회 차이 {cf['diff']:,}원 인지, 원인은 회사 소명 대기"
        return "미발견", "-", "표본·조회 모두 해당 없음"
    if i in ("A5", "A10"):
        if diff:
            return "징후", "4000C-3", f"조회 차이 {cf['diff']:,}원 인지, 원인은 회사 소명 대기"
        return "미발견", "-", "원장 금액과 맞는 세금계산서가 있어 표본 대조로는 정상으로 보임" if i == "A10" else "조회 대상 아님"
    if i in ("A6", "A7"):
        if c in cut:
            return "발견", "4000P-5", "Cut-off 테스트에서 출고일과 회계기간 불일치 확인"
        if diff:
            return "징후", "4000C-3", f"조회 차이 {cf['diff']:,}원 인지, 원인은 회사 소명 대기"
        return "미발견", "-", "Cut-off 기준금액(1천만원) 미만이라 테스트 대상에서 빠짐" if i == "A7" else "Cut-off 대상 아님"
    if i == "A8":
        return "미발견", "4000P-5", "12/26 출고 사실은 있어 Cut-off는 '일치'. 1월 반품은 후속기간 수정세금계산서를 따로 보지 않으면 안 보임"
    if i == "A9":
        if c in sample_cust:
            return "징후", "4000P-10", "표본 입금 추적 중 입금자명이 달라 추가 확인 필요"
        return "미발견", "-", "입금자명까지 대조하는 절차 없음"
    if i == "A11":
        return "발견 가능", "4000C-4", "회사 제시 연령분석표를 검토하면 드러남 (본 데모 As-Is에는 4000C-4 미포함)"


def tobe_result(a):
    hits = [x for x in tobe["exceptions"] if x["cust"] == a["cust"] and x["kind"].startswith(KT_KIND[a["id"]])]
    return hits


interval = asis["coverage"]["pop_amt"] / len(asis["sample"])
VOUCH_AMT = {"A1": None, "A2": None, "A4": 27_830_000, "A9": None}  # 표본 대조로 드러날 수 있는 건 (None=정답지 금액)


def mus_prob(a):
    """금액단위표본추출에서 이 전표가 표본에 뽑힐 확률 = 전표 금액 / 표본간격 (최대 100%)."""
    if a["id"] not in VOUCH_AMT:
        return None
    amt = VOUCH_AMT[a["id"]] or a["amount"]
    return min(1.0, amt / interval)


rows, used = [], set()
for a in truth["anomalies"]:
    ar, proc, why = asis_result(a)
    hits = tobe_result(a)
    used |= {h["no"] for h in hits}
    rows.append({"id": a["id"], "type": a["type"], "assertion": a["assertion"], "desc": a["desc"], "amount": a["amount"],
                 "asis": ar, "asis_proc": proc, "asis_why": why, "mus_prob": mus_prob(a),
                 "tobe": "발견" if hits else "미발견", "tobe_no": ", ".join(h["no"] for h in hits),
                 "tobe_detail": hits[0]["detail"] if hits else ""})
false_pos = [x for x in tobe["exceptions"] if x["no"] not in used]

s = tobe["summary"]
cov = asis["coverage"]
metrics = {
    "anomalies": rows,
    "mus_interval": interval,
    "false_positives": false_pos,
    "asis_counts": {k: sum(1 for r in rows if r["asis"] == k) for k in ("발견", "발견 가능", "징후", "미발견")},
    "tobe_counts": {k: sum(1 for r in rows if r["tobe"] == k) for k in ("발견", "미발견")},
    "coverage": {
        "asis_vouch_n": cov["vouch_n"], "asis_vouch_amt": cov["vouch_amt"],
        "pop_n": cov["pop_n"], "pop_amt": cov["pop_amt"],
        "asis_cutoff_n": cov["cutoff_n"], "tobe_cutoff_n": s["cutoff_rows"],
        "asis_conf_n": cov["conf_n"], "asis_conf_amt": cov["conf_amt"], "ar": cov["ar"],
        "tobe_matched_n": s["matched"], "tobe_matched_amt": s["matched_amount"],
        "tobe_3way_n": s["matched_paid"], "tobe_3way_amt": s["matched_paid_amount"],
        "invoices": s["invoices"], "deposits": s["deposits"],
    },
    "asis_effort": asis["effort"],
    "tobe_effort": {
        "자료 3종 업로드·실행": 1,
        "검토 필요 항목 확인 (건별 판단·회사 확인)": len(tobe["exceptions"]),
        "추정 배분·미배정 입금 확인": tobe["summary"]["rules"].get("입금-채권: 추정 배분(오래된 채권부터)", 0)
                                   + tobe["summary"]["rules"].get("입금-채권: 미배정", 0),
        "자동 생성 조서 검토·서명": 6,
    },
    "rules": tobe["summary"]["rules"],
    "vat": tobe["vat"],
    "asis_vat_diff": asis["vat_diff"],
    "confirm": [{**c, "name": next((p["name"] for p in tobe["per_customer"] if p["cust"] == c["cust"]), c["cust"]),
                 "explain": next((p["reasons"] for p in tobe["per_customer"] if p["cust"] == c["cust"]), []),
                 "evidence": next((p["evidence"] for p in tobe["per_customer"] if p["cust"] == c["cust"]), None)}
                for c in asis["confirm"]],
    "exceptions": tobe["exceptions"],
}


# 데모 페이지용 원본 미리보기 (머리글 + 앞 몇 줄)
def preview(path, sheet=None, head_row=1, n=4, cols=None):
    ws = load_workbook(path)[sheet] if sheet else load_workbook(path).active
    rows_ = list(ws.iter_rows(min_row=head_row, max_row=head_row + n, values_only=True))
    top = [[c for c in r if c is not None] for r in ws.iter_rows(min_row=1, max_row=head_row - 1, values_only=True)]
    fmt = lambda v: f"{v:,}" if isinstance(v, (int, float)) and not isinstance(v, bool) else ("" if v is None else str(v))
    body = [[fmt(v) for v in (r[:cols] if cols else r)] for r in rows_]
    return {"meta": [" · ".join(map(fmt, t)) for t in top if t], "head": body[0], "rows": body[1:],
            "total_rows": ws.max_row - head_row}


C = COMPANY_DIR
metrics["previews"] = {
    "tb": preview(C / "01_합계잔액시산표_2025.xlsx", head_row=6, n=6),
    "journal": preview(C / "02_분개장_2025.xlsx", head_row=5, n=6),
    "ledger": preview(C / "04_계정별원장_외상매출금_2025.xlsx", head_row=5, n=6),
    "subledger": preview(C / "05_거래처원장_외상매출금_2025.xlsx", head_row=5, n=5),
    "hometax": preview(C / "07_매출전자세금계산서목록_2025.xlsx", "세금계산서", head_row=6, n=4),
    "bank": preview(C / "09_은행거래내역_KB국민_보통예금.xlsx", head_row=6, n=6),
    "vatret": preview(C / "10_부가가치세신고서_과세표준요약_2025.xlsx", head_row=3, n=4),
}
(OUT / "metrics.json").write_text(json.dumps(metrics, ensure_ascii=False, indent=1, default=str))

# 엑셀 비교표
wb = Workbook()
ws = wb.active
ws.title = "이상거래_발견비교"
ws["A1"] = "심어 둔 이상 거래 11건: 지금 방식(표본) vs K-Tieout(전수)"
ws["A1"].font = BOLD
header_row(ws, 3, ["ID", "유형", "경영진 주장", "내용", "금액(원)", "As-Is 결과", "As-Is 근거 조서", "As-Is 설명",
                   "K-Tieout 결과", "K-Tieout No", "K-Tieout 설명"])
fill = {"발견": OK_FILL, "발견 가능": OK_FILL, "징후": WARN_FILL, "미발견": FLAG_FILL}
for k, r in enumerate(rows, 4):
    body_row(ws, k, [r["id"], r["type"], r["assertion"], r["desc"], r["amount"], r["asis"], r["asis_proc"], r["asis_why"],
                     r["tobe"], r["tobe_no"], r["tobe_detail"]])
    ws.cell(k, 6).fill = fill[r["asis"]]
    ws.cell(k, 9).fill = fill[r["tobe"]]
k = 4 + len(rows) + 1
ws.cell(k, 1, f"K-Tieout 오탐(정답지에 없는 검토 항목): {len(false_pos)}건").font = BOLD
widths(ws, [5, 24, 14, 50, 13, 10, 10, 50, 10, 10, 60])

ws = wb.create_sheet("커버리지")
c = metrics["coverage"]
header_row(ws, 1, ["항목", "As-Is (표본)", "K-Tieout (전수)"])
for k, (a, b, d) in enumerate([
    ("매출 증빙 대조 건수", f"{c['asis_vouch_n']} / {c['pop_n']}건 ({c['asis_vouch_n'] / c['pop_n']:.1%})",
     f"{c['pop_n']} / {c['pop_n']}건 (100%)"),
    ("매출 증빙 대조 금액", f"{c['asis_vouch_amt']:,}원 ({c['asis_vouch_amt'] / c['pop_amt']:.1%})",
     f"{c['pop_amt']:,}원 (100%) 중 3방향 일치 {c['tobe_3way_amt'] / c['pop_amt']:.1%}"),
    ("Cut-off 검토 건수", f"{c['asis_cutoff_n']}건 (1천만원 이상, 전후 5영업일)", f"{c['tobe_cutoff_n']}건 (전후 한 달 전 건)"),
    ("채권 잔액 확인", f"조회 {c['asis_conf_n']}개 거래처 ({c['asis_conf_amt'] / c['ar']:.1%})",
     "전 거래처 증빙기준 잔액 산출 (외부조회는 별도로 필요)"),
], 2):
    body_row(ws, k, [a, b, d])
widths(ws, [24, 50, 60])

ws = wb.create_sheet("작업량")
header_row(ws, 1, ["구분", "작업", "건수"])
k = 2
for lab, eff in (("As-Is", metrics["asis_effort"]), ("K-Tieout", metrics["tobe_effort"])):
    for t, n in eff.items():
        body_row(ws, k, [lab, t, n]); k += 1
ws.cell(k + 1, 1, "건당 소요시간은 파일럿에서 실측해야 합니다. 데모 페이지에서 가정치를 바꿔 볼 수 있습니다.")
widths(ws, [10, 50, 8])
wb.save(OUT / "비교_As-Is_vs_K-Tieout.xlsx")

print(f"{'ID':<4} {'유형':<22} {'As-Is':<8} {'K-Tieout':<8}")
for r in rows:
    print(f"{r['id']:<4} {r['type']:<22} {r['asis']:<8} {r['tobe']:<8} {r['tobe_no']}")
print("오탐:", [x["no"] + " " + x["kind"] for x in false_pos])
print("As-Is:", metrics["asis_counts"], "/ K-Tieout:", metrics["tobe_counts"])
