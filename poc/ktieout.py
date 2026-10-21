"""K-Tieout: 원장 × 홈택스 전자세금계산서 × 은행 거래내역 3방향 전수대사 엔진 (데모).

회사가 제출한 파일(data/company/)만 읽는다. 정답지(truth.json)는 읽지 않는다.
결과: output/to-be/K-Tieout_결과_샘플제조_2025.xlsx, output/to-be/ktieout_result.json
"""
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from itertools import combinations

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from common import (BAND_FILL, BOLD, COMPANY, COMPANY_DIR, FLAG_FILL, FY_END, FY_START, OK_FILL, OUT,
                    PERFORMANCE_MAT, SUB_END, TRIVIAL, WARN_FILL, body_row, d2s, header_row, widths, wp_header)

FEE_TOL = 3_000          # 송금수수료 차감 허용폭
DATE_WINDOW = 20         # 원장-세금계산서 일자 허용 차이(일)
AGING_DAYS = 150         # 장기 미회수 판단 기준(작성일~기말)
CUTOFF_FROM, CUTOFF_TO = date(2025, 12, 15), date(2026, 1, 15)


def P(s):
    return datetime.strptime(s, "%Y-%m-%d").date()


# ---------------------------------------------------------------- 1. 읽기 (정규화)
def read_customers():
    ws = load_workbook(COMPANY_DIR / "06_거래처등록현황.xlsx").active
    out = {}
    for r in ws.iter_rows(min_row=4, values_only=True):
        if r[0]:
            out[r[0]] = {"code": r[0], "name": r[1], "bizno": r[2]}
    return out


def read_journal(fname):
    ws = load_workbook(COMPANY_DIR / fname).active
    rows = []
    for r in ws.iter_rows(min_row=6, values_only=True):
        if not r[0]:
            continue
        rows.append({"date": P(r[0]), "no": r[1], "acct": r[3], "cust": r[5], "memo": r[7],
                     "dr": r[8] or 0, "cr": r[9] or 0})
    return rows


def read_hometax(fname):
    ws = load_workbook(COMPANY_DIR / fname)["세금계산서"]
    head = [c.value for c in ws[6]]
    assert head[:3] == ["작성일자", "승인번호", "발급일자"], "홈택스 목록 형식이 아닙니다"
    out = []
    for r in ws.iter_rows(min_row=7, values_only=True):
        if not r[1]:
            continue
        orig = re.search(r"당초승인번호 ([0-9a-f-]+)", r[20] or "")
        out.append({"appr": r[1], "wdate": P(r[0]), "idate": P(r[2]), "bizno": r[9], "name": r[11],
                    "total": r[14], "supply": r[15], "vat": r[16], "kind": r[17], "orig": orig.group(1) if orig else None,
                    "item": r[26]})
    return out


def read_bank():
    ws = load_workbook(COMPANY_DIR / "09_은행거래내역_KB국민_보통예금.xlsx").active
    out = []
    for i, r in enumerate(ws.iter_rows(min_row=7, values_only=True)):
        if not r[0]:
            continue
        out.append({"id": f"B{i + 1:04d}", "dt": datetime.strptime(r[0], "%Y.%m.%d %H:%M:%S"), "kind": r[1],
                    "memo": r[2], "out": r[3] or 0, "in": r[4] or 0})
    return out


def read_opening():
    ws = load_workbook(COMPANY_DIR / "05_거래처원장_외상매출금_2025.xlsx").active
    op, bal = {}, {}
    for r in ws.iter_rows(min_row=6, values_only=True):
        if r[0]:
            op[r[0]] = r[4] or 0
            bal[r[0]] = r[7] or 0
    return op, bal


def read_vat_return():
    ws = load_workbook(COMPANY_DIR / "10_부가가치세신고서_과세표준요약_2025.xlsx").active
    return [r[6] for r in ws.iter_rows(min_row=4, values_only=True) if r[0]]


def norm(s):
    s = re.sub(r"\(주\)|주\)|\(주|㈜|주식회사|\s", "", s or "")
    return s


# ---------------------------------------------------------------- 2. 대사
def run():
    custs = read_customers()
    biz2code = {c["bizno"]: k for k, c in custs.items()}
    j = read_journal("02_분개장_2025.xlsx") + read_journal("03_분개장_2026_1-2월.xlsx")
    inv = read_hometax("07_매출전자세금계산서목록_2025.xlsx") + read_hometax("08_매출전자세금계산서목록_2026_1-2월.xlsx")
    bank = read_bank()
    opening, ledger_bal = read_opening()
    vat_ret = read_vat_return()

    # 원장 매출: 외상매출금 차변 (분개장 한 전표의 108 차변 + 404 대변)
    sales = []
    for r in j:
        if r["acct"] == "108" and r["dr"]:
            sales.append({"id": f"L{len(sales) + 1:04d}", "date": r["date"], "no": r["no"], "cust": r["cust"],
                          "total": r["dr"], "memo": r["memo"], "inv": None, "rule": None})
    supply_by_key = defaultdict(list)
    for r in j:
        if r["acct"] == "404":
            supply_by_key[(r["date"], r["no"])].append(r["cr"])
    for s in sales:
        s["supply"] = (supply_by_key.get((s["date"], s["no"])) or [round(s["total"] / 1.1)])[0]
    coll = [r for r in j if r["acct"] == "108" and r["cr"]]
    suspense = [r for r in j if r["acct"] == "257" and r["cr"]]

    for i in inv:
        i["cust"] = biz2code.get(i["bizno"])
        i["ledger"] = None
        i["rule"] = None

    log = Counter()
    # (1) 세금계산서 ↔ 원장: 같은 거래처·같은 금액, 일자가 가장 가까운 전표
    for i in sorted(inv, key=lambda x: x["wdate"]):
        cands = [s for s in sales if s["inv"] is None and s["cust"] == i["cust"] and s["total"] == i["total"]
                 and abs((s["date"] - i["wdate"]).days) <= DATE_WINDOW]
        if not cands:
            continue
        s = min(cands, key=lambda s: abs((s["date"] - i["wdate"]).days))
        diff = (s["date"] - i["wdate"]).days
        s["inv"], i["ledger"] = i["appr"], s["id"]
        s["rule"] = i["rule"] = "일자·금액 일치" if diff == 0 else f"금액 일치·일자 {diff:+d}일"
        log["원장-세금계산서: 일자·금액 일치" if diff == 0 else "원장-세금계산서: 금액 일치, 일자 차이"] += 1

    exceptions = []

    def X(kind, assertion, risk, amount, detail, action, ledger=None, appr=None, cust=None, deposit=None):
        exceptions.append({"no": f"EX-{len(exceptions) + 1:02d}", "kind": kind, "assertion": assertion, "risk": risk,
                           "amount": amount, "detail": detail, "action": action, "ledger": ledger, "appr": appr,
                           "cust": cust, "cust_name": custs.get(cust, {}).get("name", ""), "deposit": deposit})

    # (2) 남은 것끼리 금액 불일치 쌍 찾기 (같은 거래처, 일자 ±5일, 금액 ±20%)
    for i in inv:
        if i["ledger"] or i["total"] <= 0:
            continue
        cands = [s for s in sales if s["inv"] is None and s["cust"] == i["cust"]
                 and abs((s["date"] - i["wdate"]).days) <= 5 and abs(s["total"] - i["total"]) <= 0.2 * i["total"]]
        if cands:
            s = min(cands, key=lambda s: abs(s["total"] - i["total"]))
            s["inv"], i["ledger"] = i["appr"], s["id"]
            s["rule"] = i["rule"] = "금액 불일치"
            log["원장-세금계산서: 금액 불일치(쌍)"] += 1
            if s["date"].year == 2025:
                X("금액 불일치", "정확성", "높음" if abs(s["total"] - i["total"]) >= TRIVIAL else "중간",
                  s["total"] - i["total"],
                  f"원장 {s['total']:,} vs 세금계산서 {i['total']:,} (공급가액 {s['supply']:,} vs {i['supply']:,})",
                  "회사에 정정 근거 요청, 수정분개 필요 여부 판단", ledger=s["id"], appr=i["appr"], cust=s["cust"])

    # (3) 원장에만 있는 매출
    for s in sales:
        if s["inv"] or s["date"].year != 2025:
            continue
        twin = [t for t in sales if t is not s and t["inv"] and t["cust"] == s["cust"] and t["date"] == s["date"]
                and t["total"] == s["total"]]
        if twin:
            s["rule"] = "중복 의심"
            X("원장 중복 계상 의심", "실재성", "높음", s["total"],
              f"{d2s(s['date'])} 같은 거래처·같은 금액 전표가 2건, 세금계산서는 1건({twin[0]['inv']})",
              "중복 전표 여부 확인, 수정분개 필요 여부 판단", ledger=s["id"], appr=twin[0]["inv"], cust=s["cust"])
        else:
            s["rule"] = "세금계산서 없음"
    # (4) 세금계산서에만 있는 것
    for i in inv:
        if i["ledger"]:
            continue
        if i["total"] < 0:
            o = next((x for x in inv if x["appr"] == i["orig"]), None)
            i["rule"] = "수정세금계산서 원장 미반영"
            if i["wdate"].year == 2025:
                X("수정세금계산서 원장 미반영", "정확성·실재성", "높음", i["total"],
                  f"{d2s(i['wdate'])} {i['kind']} {i['total']:,}원 (당초 {i['orig']}), 원장에 매출 차감 없음",
                  "반품·에누리 사실 확인, 매출·채권 차감 분개 필요 여부 판단", appr=i["appr"], cust=i["cust"])
        elif i["wdate"].year == 2025:
            i["rule"] = "원장 매출 없음"
            near = i["wdate"] >= CUTOFF_FROM
            X("세금계산서O·원장X" + (" (기말 근접)" if near else ""), "완전성" + ("·기간귀속" if near else ""), "높음",
              -i["total"], f"{d2s(i['wdate'])} 발행 {i['total']:,}원, 2025년·2026년 원장 어디에도 매출 없음",
              "매출 누락 여부 확인 (입금 처리 계정 확인)", appr=i["appr"], cust=i["cust"])

    # (5) 기간귀속: 연도가 다른 원장·세금계산서 쌍
    cutoff_rows = []
    for s in sales:
        if s["inv"] is None:
            continue
        i = next(x for x in inv if x["appr"] == s["inv"])
        in_window = CUTOFF_FROM <= s["date"] <= CUTOFF_TO or CUTOFF_FROM <= i["wdate"] <= CUTOFF_TO
        if in_window:
            cutoff_rows.append((s, i))
        if s["date"].year != i["wdate"].year:
            over = s["date"].year == 2025
            X("기간귀속 " + ("(당기 과대)" if over else "(당기 과소)"), "기간귀속", "높음",
              (1 if over else -1) * s["total"],
              f"원장 {d2s(s['date'])} / 세금계산서 작성일 {d2s(i['wdate'])}",
              "출고증·인수증으로 실제 인도일 확인, 귀속연도 판단", ledger=s["id"], appr=i["appr"], cust=s["cust"])
    # (6) 기말 직전 매출의 후속기간 취소
    for i in inv:
        if i["total"] < 0 and i["wdate"] > FY_END and i["orig"]:
            o = next((x for x in inv if x["appr"] == i["orig"]), None)
            if o and o["wdate"] >= CUTOFF_FROM and o["wdate"] <= FY_END:
                X("기말 직전 매출 후속 취소", "실재성·기간귀속", "높음", o["total"],
                  f"{d2s(o['wdate'])} 매출 {o['total']:,}원이 {d2s(i['wdate'])} 수정세금계산서로 {i['total']:,}원 취소",
                  "반품 사유·물류 기록 확인, 당기 매출 인식 적정성 판단", appr=o["appr"], cust=o["cust"])

    # (7) 은행 입금 → 거래처
    by_norm = {norm(c["name"]): k for k, c in custs.items()}

    def who(memo):
        n = norm(memo)
        if not n:
            return None
        hits = {k for nn, k in by_norm.items() if nn == n or (len(n) >= 3 and (nn.startswith(n) or n.startswith(nn)))}
        return hits.pop() if len(hits) == 1 else None

    deps = [b for b in bank if b["in"]]
    for d in deps:
        d["cust"] = who(d["memo"]) if d["kind"] != "이자" else None
        d["alloc"] = []
        d["rule"] = None

    # 받을 채권 항목: 전기이월 + 세금계산서(수정분은 당초에 합산) + 세금계산서 없는 원장 매출
    items = []
    for code, amt in opening.items():
        if amt:
            items.append({"key": f"OPEN-{code}", "cust": code, "date": FY_START - timedelta(days=1), "amt": amt,
                          "src": "전기이월"})
    net = {i["appr"]: i["total"] for i in inv if i["total"] > 0}
    net_ye = dict(net)  # 기말 기준: 기말 이후 수정분은 빼지 않는다
    for i in inv:
        if i["total"] < 0 and i["orig"] in net:
            net[i["orig"]] += i["total"]
            if i["wdate"] <= FY_END:
                net_ye[i["orig"]] += i["total"]
    for i in inv:
        if i["total"] > 0 and net_ye[i["appr"]] > 0:
            items.append({"key": i["appr"], "cust": i["cust"], "date": i["wdate"], "amt": net[i["appr"]],
                          "amt_ye": net_ye[i["appr"]], "src": "세금계산서"})
    for s in sales:
        if s["rule"] == "세금계산서 없음":
            items.append({"key": s["id"], "cust": s["cust"], "date": s["date"], "amt": s["total"], "src": "원장(증빙없음)"})
    for it in items:
        it["paid"] = 0
        it["paid_by_ye"] = 0
        it["deps"] = []

    def assign(d, its, rule, partial=None):
        for it in its:
            take = partial if partial is not None else it["amt"] - it["paid"]
            it["paid"] += take
            if d["dt"].date() <= FY_END:
                it["paid_by_ye"] += take
            it["deps"].append(d["id"])
            d["alloc"].append((it["key"], take))
        d["rule"] = rule
        log["입금-채권: " + rule] += 1

    by_c = defaultdict(list)
    for it in items:
        by_c[it["cust"]].append(it)
    rem = lambda it: it["amt"] - it["paid"]

    def open_items(cust, d):
        pool = by_c[cust] if cust else items
        return sorted([it for it in pool if rem(it) > 0 and it["date"] <= d["dt"].date()], key=lambda it: it["date"])

    # 1단계: 입금 1건 ↔ 채권 1건 또는 여러 건 (단건·수수료 차감·합산)
    for d in sorted(deps, key=lambda d: d["dt"]):
        if not d["cust"]:
            continue
        a, open_ = d["in"], open_items(d["cust"], d)
        hit = next((it for it in open_ if rem(it) == a), None)
        if hit:
            assign(d, [hit], "단건 일치"); continue
        hit = next((it for it in open_ if 0 < rem(it) - a <= FEE_TOL), None)
        if hit:
            assign(d, [hit], "수수료 차감 입금"); continue
        grp = next((open_[st:st + k] for k in range(2, 7) for st in range(max(0, len(open_) - k + 1))
                    if 0 <= sum(map(rem, open_[st:st + k])) - a <= FEE_TOL), None)
        if grp:
            assign(d, grp, "합산 입금"); continue
        grp = next((list(g) for k in (2, 3) for g in combinations(open_[:12], k)
                    if 0 <= sum(map(rem, g)) - a <= FEE_TOL), None)
        if grp:
            assign(d, grp, "합산 입금(비연속)")

    # 2단계: 남은 입금 2~3건의 합 ↔ 채권 1건 (분할 입금). 거래처 미식별 입금은 같은 입금자명끼리 묶는다
    def split_pass(cust_of_dep, label):
        left = defaultdict(list)
        for d in deps:
            if not d["alloc"] and d["kind"] != "이자" and cust_of_dep(d) is not None:
                left[cust_of_dep(d)].append(d)
        hits = []
        for key, ds in left.items():
            ds.sort(key=lambda d: d["dt"])
            pool = by_c[key] if label == "분할 입금" else items
            for it in sorted([it for it in pool if rem(it) > 0], key=lambda it: it["date"]):
                cand = [d for d in ds if not d["alloc"] and it["date"] <= d["dt"].date() <= it["date"] + timedelta(days=200)]
                g = next((g for k in (2, 3) for g in combinations(cand[:10], k)
                          if 0 <= rem(it) - sum(d["in"] for d in g) <= FEE_TOL), None)
                if g:
                    for d in g:
                        assign(d, [it], label, partial=d["in"])
                    hits.append((it, g))
        return hits

    split_pass(lambda d: d["cust"], "분할 입금")

    # 3단계: 그래도 남은 거래처 입금은 오래된 채권부터 추정 배분
    for d in sorted(deps, key=lambda d: d["dt"]):
        if not d["cust"] or d["alloc"]:
            continue
        a = d["in"]
        for it in open_items(d["cust"], d):
            if a <= 0:
                break
            take = min(a, rem(it))
            assign(d, [it], "추정 배분(오래된 채권부터)", partial=take)
            a -= take
        if not d["alloc"]:
            d["rule"] = "미배정(거래처 확인됨)"
            log["입금-채권: 미배정"] += 1

    # 거래처를 못 찾은 입금 → 열린 채권과 금액(단건 또는 같은 입금자 분할 합계)이 같으면 제3자 입금 의심
    def dep_txt(d):
        return f"{d['dt']:%Y-%m-%d} {d['in']:,}"

    def third(it, ds):
        X("제3자 명의 입금", "실재성(거래 상대방)", "중간", sum(d["in"] for d in ds),
          f"입금자 '{ds[0]['memo']}' {len(ds)}건 ({', '.join(dep_txt(d) for d in ds)})"
          f" → {custs[it['cust']]['name']} 채권({it['key']})과 금액 일치",
          "입금자와 매출처 관계 확인 (특수관계자 여부 포함)", appr=it["key"], cust=it["cust"],
          deposit=", ".join(d["id"] for d in ds))

    for d in deps:
        if d["kind"] == "이자":
            d["rule"] = "비매출(이자)"
            continue
        if d["cust"] or d["alloc"]:
            continue
        hit = [it for it in items if rem(it) == d["in"] and it["date"] <= d["dt"].date()]
        if len(hit) == 1:
            assign(d, hit, "제3자 입금 의심")
            third(hit[0], [d])
    for it, g in split_pass(lambda d: None if d["cust"] or d["alloc"] else d["memo"], "제3자 입금 의심"):
        third(it, list(g))
    for d in deps:
        if not d["alloc"] and d["kind"] != "이자" and not d["cust"]:
            d["rule"] = "미확인 입금"
            log["입금-채권: 미확인 입금"] += 1

    # 원장에만 있는 매출의 입금 여부
    for s in sales:
        if s["rule"] != "세금계산서 없음":
            continue
        it = next(it for it in items if it["key"] == s["id"])
        if it["paid"]:
            X("원장O·세금계산서X·입금O", "발생(세무)", "중간", s["total"],
              f"{d2s(s['date'])} {s['total']:,}원 매출·입금({', '.join(it['deps'])}) 있으나 세금계산서 없음",
              "세금계산서 미발행 사유 확인, 부가세 신고 누락 여부 확인", ledger=s["id"], cust=s["cust"])
        else:
            X("원장O·세금계산서X·입금X", "실재성", "매우 높음", s["total"],
              f"{d2s(s['date'])} {s['total']:,}원, 세금계산서·입금(2026-02-28까지) 모두 없음",
              "계약서·출고증·거래처 실체 확인, 외부조회 우선 대상", ledger=s["id"], cust=s["cust"])

    # 세금계산서O·원장X 건이 입금됐다면 어느 계정으로 처리됐는지 단서 추가
    for x in exceptions:
        if x["kind"].startswith("세금계산서O·원장X"):
            it = next((it for it in items if it["key"] == x["appr"]), None)
            if it and it["deps"]:
                d = next(d for d in deps if d["id"] == it["deps"][0])
                sus = [r for r in suspense if r["date"] == d["dt"].date()]
                x["detail"] += f" / 입금 {d['dt']:%Y-%m-%d} {d['in']:,}원" + (" → 원장상 가수금(257) 처리" if sus else "")

    # (8) 장기 미회수
    for it in items:
        if it["src"] == "세금계산서" and it["date"].year == 2025 and it["paid"] == 0 \
                and (FY_END - it["date"]).days > AGING_DAYS:
            X("장기 미회수", "평가", "중간", it["amt"],
              f"작성일 {d2s(it['date'])}, 기말까지 {(FY_END - it['date']).days}일 경과, 2026-02-28까지 입금 없음",
              "회수 가능성 검토, 대손충당금 적정성 판단", appr=it["key"], cust=it["cust"])

    # (9) 거래처별 기말 채권: 장부 vs 증빙 기준(세금계산서·입금)
    per = []
    for code in sorted(set(ledger_bal) | {it["cust"] for it in items}):
        ev = 0
        for it in items:
            if it["cust"] != code or it["date"] > FY_END or it["src"] == "원장(증빙없음)":
                continue
            left = it.get("amt_ye", it["amt"]) - it["paid_by_ye"]
            ev += 0 if it["paid_by_ye"] and left <= FEE_TOL else left  # 수수료 차감분은 회수로 본다
        lb = ledger_bal.get(code, 0)
        reasons = [x["no"] for x in exceptions if x["cust"] == code and x["kind"] != "장기 미회수"]
        per.append({"cust": code, "name": custs[code]["name"], "ledger": lb, "evidence": ev, "diff": lb - ev,
                    "reasons": reasons})

    # (10) 부가세 대사 (분기)
    q = lambda d: (d.month - 1) // 3
    ht_q = [sum(i["supply"] for i in inv if i["wdate"].year == 2025 and q(i["wdate"]) == k) for k in range(4)]
    gl_q = [sum(s["supply"] for s in sales if s["date"].year == 2025 and q(s["date"]) == k) for k in range(4)]

    return {"custs": custs, "sales": sales, "inv": inv, "deps": deps, "items": items, "exceptions": exceptions,
            "per": per, "cutoff": cutoff_rows, "ht_q": ht_q, "gl_q": gl_q, "vat_ret": vat_ret, "log": log,
            "bank": bank, "coll": coll}


# ---------------------------------------------------------------- 3. 조서 출력
def write(res):
    out = OUT / "to-be"
    out.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    custs = res["custs"]
    s25 = [s for s in res["sales"] if s["date"].year == 2025]
    i25 = [i for i in res["inv"] if i["wdate"].year == 2025]
    ex = res["exceptions"]
    matched = [s for s in s25 if s["inv"] and s["rule"] != "금액 불일치"]
    matched_paid = []
    for s in matched:
        it = next((it for it in res["items"] if it["key"] == s["inv"]), None)
        if it and it["paid"] >= it["amt"] - FEE_TOL:
            matched_paid.append(s)

    # 00 요약
    ws = wb.active
    ws.title = "00_요약"
    r = wp_header(ws, "KT-00", "K-Tieout 3방향 전수대사 요약 (매출·매출채권)", "K-Tieout(자동)", "(검토자 서명란)",
                  datetime.now().strftime("%Y.%m.%d"), "", "도구는 장부를 고치지 않습니다. 아래 결과는 회계사 판단을 위한 자료입니다.")
    rows = [
        ("입력 자료", ""),
        ("원장 매출 전표 (2025, 외상매출금 차변)", f"{len(s25):,}건 / {sum(s['total'] for s in s25):,}원"),
        ("홈택스 매출 전자세금계산서 (작성일 2025)", f"{len(i25):,}건 / {sum(i['total'] for i in i25):,}원"),
        ("은행 입금 (2025.01~2026.02)", f"{len(res['deps']):,}건 / {sum(d['in'] for d in res['deps']):,}원"),
        ("", ""),
        ("대사 결과", ""),
        ("원장↔세금계산서 일치 (금액 기준)", f"{len(matched):,}건 ({len(matched) / len(s25):.1%}) / {sum(s['total'] for s in matched):,}원"),
        ("  └ 입금까지 확인 (3방향 일치, 2026-02-28 기준)", f"{len(matched_paid):,}건 / {sum(s['total'] for s in matched_paid):,}원"),
        ("검토 필요 항목", f"{len(ex)}건 (아래 01 시트)"),
        ("", ""),
        ("검토 필요 항목 유형별", ""),
    ] + [(f"  {k}", f"{v}건") for k, v in Counter(x["kind"] for x in ex).most_common()] + [
        ("", ""),
        ("매칭 규칙별 처리 건수", ""),
    ] + [(f"  {k}", f"{v:,}건") for k, v in sorted(res["log"].items())]
    for k, v in rows:
        ws.cell(r, 1, k).font = BOLD if v == "" else Font()
        ws.cell(r, 2, v)
        r += 1
    widths(ws, [52, 60])

    # 01 검토 필요 목록
    ws = wb.create_sheet("01_검토필요목록")
    r = wp_header(ws, "KT-01", "검토 필요 항목 (회계사 판단란 포함)", "K-Tieout(자동)", "", "", "")
    header_row(ws, r, ["No", "유형", "관련 경영진 주장", "위험", "거래처", "금액(원)", "내용", "권고 확인 절차",
                       "원장 ID", "세금계산서 승인번호", "입금 ID", "회계사 판단", "조치 결과", "Tick"])
    for k, x in enumerate(ex, r + 1):
        body_row(ws, k, [x["no"], x["kind"], x["assertion"], x["risk"], x["cust_name"], x["amount"], x["detail"],
                         x["action"], x["ledger"], x["appr"], x["deposit"], "", "", ""],
                 fill=FLAG_FILL if x["risk"] in ("높음", "매우 높음") else WARN_FILL)
    widths(ws, [7, 26, 14, 8, 16, 14, 70, 44, 8, 30, 8, 18, 18, 6])
    ws.freeze_panes = ws.cell(r + 1, 1)

    # 02 기간귀속 (4000P-5 형식, 대상 전수)
    ws = wb.create_sheet("02_매출Cut-off")
    r = wp_header(ws, "4000P-5", "매출Cut-off (자동 전수: 2025-12-15 ~ 2026-01-15)", "K-Tieout(자동)", "", "", "",
                  "· 테스트 절차/방법: 결산일 전후 한 달의 원장 매출 전 건을 세금계산서 작성일과 대조")
    header_row(ws, r, ["No.", "전표일", "전표번호", "내역", "거래처", "금액", "첨부증빙(승인번호)", "증빙일(세금계산서 작성일)",
                       "회계기간 일치여부", "비고"])
    for k, (s, i) in enumerate(sorted(res["cutoff"], key=lambda p: p[0]["date"]), 1):
        ok = s["date"].year == i["wdate"].year
        body_row(ws, r + k, [k, d2s(s["date"]), s["no"], s["memo"], custs[s["cust"]]["name"], s["total"], i["appr"],
                             d2s(i["wdate"]), "일치" if ok else "불일치", "" if ok else "검토필요목록 참조"],
                 fill=None if ok else FLAG_FILL)
    widths(ws, [5, 11, 8, 30, 16, 14, 30, 14, 10, 18])

    # 03 부가세대사표 (4000P-2 형식)
    ws = wb.create_sheet("03_부가세대사표")
    r = wp_header(ws, "4000P-2", "부가세대사표 (자동, 차이 원인 분해 포함)", "K-Tieout(자동)", "", "", "")
    header_row(ws, r, ["공급가액", "1/4분기", "2/4분기", "3/4분기", "4/4분기", "계"])
    ret = res["vat_ret"]
    lines = [("(매출 부가세 신고) 과세(세금계산서)", ret), ("홈택스 발급분 합계 (작성일 기준)", res["ht_q"]),
             ("손익계산서 매출액 (원장 404)", res["gl_q"]),
             ("차이 (원장 − 신고)", [g - v for g, v in zip(res["gl_q"], ret)])]
    for k, (lab, vals) in enumerate(lines, r + 1):
        body_row(ws, k, [lab] + list(vals) + [sum(vals)], fill=BAND_FILL if lab.startswith("차이") else None)
    k = r + len(lines) + 2
    ws.cell(k, 1, "차이 원인 분해 (공급가액 기준, 검토필요목록과 연결)").font = BOLD
    header_row(ws, k + 1, ["항목", "1/4분기", "2/4분기", "3/4분기", "4/4분기", "계", "관련 No"])
    q = lambda d: (d.month - 1) // 3
    explained = [0, 0, 0, 0]
    vat_items = []
    rr = k + 2
    sales_by_id = {s["id"]: s for s in res["sales"]}
    inv_by = {i["appr"]: i for i in res["inv"]}
    for x in ex:
        if x["kind"] in ("장기 미회수", "제3자 명의 입금", "기말 직전 매출 후속 취소"):
            continue
        s = sales_by_id.get(x["ledger"])
        i = inv_by.get(x["appr"])
        vals = [0, 0, 0, 0]
        if x["kind"] == "금액 불일치":
            vals[q(s["date"])] = s["supply"] - i["supply"]
        elif x["kind"].startswith("기간귀속"):
            if s["date"].year == 2025:
                vals[q(s["date"])] = s["supply"]
            else:
                vals[q(i["wdate"])] = -i["supply"]
        elif x["kind"].startswith("원장O") or x["kind"] == "원장 중복 계상 의심":
            vals[q(s["date"])] = s["supply"]
        elif x["kind"].startswith("세금계산서O") or x["kind"] == "수정세금계산서 원장 미반영":
            vals[q(i["wdate"])] = -i["supply"]
        explained = [a + b for a, b in zip(explained, vals)]
        vat_items.append({"no": x["no"], "label": f"{x['kind']} · {x['cust_name']}", "v": sum(vals)})
        body_row(ws, rr, [x["kind"] + " · " + x["cust_name"]] + vals + [sum(vals), x["no"]])
        rr += 1
    diff = [g - v for g, v in zip(res["gl_q"], ret)]
    body_row(ws, rr, ["설명된 차이 합계"] + explained + [sum(explained), ""], fill=BAND_FILL)
    body_row(ws, rr + 1, ["미설명 차이"] + [d - e for d, e in zip(diff, explained)] + [sum(diff) - sum(explained), ""],
             fill=OK_FILL if sum(diff) == sum(explained) else FLAG_FILL)
    widths(ws, [44, 16, 16, 16, 16, 18, 10])

    # 04 거래처별 채권 대사 (4000C-2 형식 + 증빙 기준 잔액)
    ws = wb.create_sheet("04_거래처별채권대사")
    r = wp_header(ws, "4000C-2", "계정내역,명세대사 + 증빙기준 잔액 (자동)", "K-Tieout(자동)", "", "", "",
                  "증빙기준 잔액 = 전기이월 + 세금계산서(수정 반영) − 2025-12-31까지 입금. 채권채무조회 회신 차이 설명에 사용")
    header_row(ws, r, ["계정", "거래처", "기말잔액(장부)", "증빙기준 잔액", "차이", "Tick-Mark", "관련 No", "비고"])
    for k, p in enumerate(res["per"], r + 1):
        body_row(ws, k, ["외상매출금", p["name"], p["ledger"], p["evidence"], p["diff"], "GL" if p["diff"] == 0 else "",
                         ", ".join(p["reasons"]), ""], fill=None if p["diff"] == 0 else WARN_FILL)
    k = r + len(res["per"]) + 1
    body_row(ws, k, ["", "합      계", sum(p["ledger"] for p in res["per"]), sum(p["evidence"] for p in res["per"]),
                     sum(p["diff"] for p in res["per"]), "F", "", ""], fill=BAND_FILL)
    widths(ws, [11, 18, 16, 16, 14, 9, 22, 20])

    # 05 전수대사표
    ws = wb.create_sheet("05_전수대사표")
    r = wp_header(ws, "KT-05", "원장 매출 전 건 대사표 (2025)", "K-Tieout(자동)", "", "", "")
    header_row(ws, r, ["원장 ID", "전표일", "전표번호", "거래처", "원장 금액", "세금계산서 승인번호", "작성일", "세금계산서 금액",
                       "입금 ID", "입금 상태(2026-02-28)", "대사 규칙", "Tick"])
    for k, s in enumerate(s25, r + 1):
        i = inv_by.get(s["inv"])
        it = next((it for it in res["items"] if it["key"] in (s["inv"], s["id"])), None)
        st = "" if not it else ("입금완료" if it["paid"] >= it["amt"] - FEE_TOL else ("부분입금" if it["paid"] else "미입금"))
        ok = i and s["rule"] != "금액 불일치" and st == "입금완료"
        body_row(ws, k, [s["id"], d2s(s["date"]), s["no"], custs[s["cust"]]["name"], s["total"], i and i["appr"],
                         i and d2s(i["wdate"]), i and i["total"], it and ", ".join(it["deps"]), st, s["rule"],
                         "√√√" if ok else ""], fill=None if i and s["rule"] != "금액 불일치" else FLAG_FILL)
    widths(ws, [8, 11, 8, 16, 14, 30, 11, 14, 22, 12, 18, 6])
    ws.freeze_panes = ws.cell(r + 1, 1)

    # 06 입금 매칭 내역
    ws = wb.create_sheet("06_입금매칭내역")
    r = wp_header(ws, "KT-06", "은행 입금 매칭 내역", "K-Tieout(자동)", "", "", "")
    header_row(ws, r, ["입금 ID", "거래일시", "기재내용", "맡기신금액", "식별 거래처", "배정 채권", "매칭 규칙"])
    for k, d in enumerate(res["deps"], r + 1):
        body_row(ws, k, [d["id"], f"{d['dt']:%Y-%m-%d %H:%M}", d["memo"], d["in"],
                         custs.get(d["cust"], {}).get("name", ""), ", ".join(a for a, _ in d["alloc"]), d["rule"]],
                 fill=FLAG_FILL if d["rule"] in ("제3자 입금 의심", "미확인 입금", "미배정(거래처 확인됨)") else None)
    widths(ws, [8, 17, 16, 15, 16, 50, 20])

    path = out / "K-Tieout_결과_샘플제조_2025.xlsx"
    wb.save(path)
    json.dump({
        "exceptions": ex,
        "summary": {"ledger_sales": len(s25), "ledger_amount": sum(s["total"] for s in s25),
                    "invoices": len(i25), "deposits": len(res["deps"]), "matched": len(matched),
                    "matched_amount": sum(s["total"] for s in matched), "matched_paid": len(matched_paid),
                    "matched_paid_amount": sum(s["total"] for s in matched_paid), "rules": dict(res["log"]),
                    "cutoff_rows": len(res["cutoff"])},
        "vat": {"return": res["vat_ret"], "hometax": res["ht_q"], "ledger": res["gl_q"], "items": vat_items},
        "per_customer": res["per"],
    }, open(out / "ktieout_result.json", "w"), ensure_ascii=False, indent=1, default=str)
    return path


if __name__ == "__main__":
    res = run()
    path = write(res)
    print("결과:", path)
    for x in res["exceptions"]:
        print(f"  {x['no']} {x['kind']:<22} {x['cust_name']:<12} {x['amount']:>14,}  {x['detail'][:70]}")
    print("규칙:", dict(res["log"]))
