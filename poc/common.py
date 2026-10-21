"""데모 공통: 회사 정보, 기간, 엑셀 서식 도우미."""
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
COMPANY_DIR = DATA / "company"
OUT = ROOT / "output"

COMPANY = {
    "name": "샘플제조(주)",
    "bizno": "123-81-45678",
    "ceo": "박정민",
    "addr": "경기도 안산시 단원구 가상산업로 100",
    "email": "tax@sample-mfg.example",
}
FY_START, FY_END = date(2025, 1, 1), date(2025, 12, 31)
SUB_END = date(2026, 2, 28)  # 후속기간(기말 후 확인 범위) 끝

# 감사계획 단계에서 정한 중요성 (가정치: 매출액의 0.5% 기준)
MATERIALITY = 50_000_000
PERFORMANCE_MAT = 37_500_000
TRIVIAL = 2_500_000

THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="DDE4EE")
BAND_FILL = PatternFill("solid", fgColor="F2F2F2")
FLAG_FILL = PatternFill("solid", fgColor="FDE2E1")
WARN_FILL = PatternFill("solid", fgColor="FFF4CC")
OK_FILL = PatternFill("solid", fgColor="E3F2E6")
BOLD = Font(bold=True)
NUM = "#,##0;-#,##0;0"


def d2s(d):
    return d.strftime("%Y-%m-%d") if d else ""


def biz_days(start, end):
    d = start
    while d <= end:
        if d.weekday() < 5:
            yield d
        d += timedelta(days=1)


def next_biz(d):
    while d.weekday() >= 5:
        d += timedelta(days=1)
    return d


def month_end(y, m):
    return (date(y + m // 12, m % 12 + 1, 1) - timedelta(days=1))


def header_row(ws, row, values, fill=HEAD_FILL):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = BOLD
        c.fill = fill
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def body_row(ws, row, values, fill=None, num_cols=()):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.border = BORDER
        if fill:
            c.fill = fill
        if i in num_cols or (isinstance(v, (int, float)) and not isinstance(v, bool)):
            c.number_format = NUM


def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


def wp_header(ws, no, title, preparer, reviewer, prep_date, rev_date, note=""):
    """감사마스터 조서 머리: 조서번호·제목·작성자·검토자."""
    ws["A1"] = f"( {no} ) {title}"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"회사명: {COMPANY['name']}    결산일: {d2s(FY_END)}    감사인: 가나회계법인(가상)"
    ws["A3"] = f"작성자 : {preparer} ( {prep_date} )    검토자 : {reviewer} ( {rev_date} )"
    if note:
        ws["A4"] = note
        ws["A4"].font = Font(italic=True, color="666666")
    return 6


TICKMARKS = [
    ("PY", "전기 감사보고서와 일치(Agreed to Prior Year Audit Report)"),
    ("PTB", "시산표와 대사(Agreed to Trial Balance)"),
    ("CY", "당기 감사보고서와 일치(Agreed to Current Year Audit Report)"),
    ("GL", "GL과 일치(Agreed to General Ledger)"),
    ("F", "합계 검증(Footed)"),
    ("Rx", "재계산(Recalculation)"),
]
