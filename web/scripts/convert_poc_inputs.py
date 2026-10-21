"""poc/data/company의 가상 엑셀 자료를 웹 엔진 입력(JSON)으로 정규화한다.

실행: poc/.venv/bin/python web/scripts/convert_poc_inputs.py
출력: web/src/data/mock/tieout/input.json, web/src/data/mock/tieout/asis.json
엔진은 회사 제출 파일만 쓴다. asis.json은 '현행 표본 방식' 비교용(정답지 기준) 시연 자료다.
"""
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "poc" / "data" / "company"
OUT = ROOT / "web" / "src" / "data" / "mock" / "tieout"


def rows(fname, min_row, sheet=None):
    wb = load_workbook(SRC / fname, read_only=True)
    ws = wb[sheet] if sheet else wb.active
    return [r for r in ws.iter_rows(min_row=min_row, values_only=True)]


customers = [{"code": r[0], "name": r[1], "bizno": r[2]} for r in rows("06_거래처등록현황.xlsx", 4) if r[0]]

journal = []
for f in ("02_분개장_2025.xlsx", "03_분개장_2026_1-2월.xlsx"):
    for r in rows(f, 6):
        if r[0] and str(r[3]) in ("108", "404", "257"):
            journal.append({"date": r[0], "no": r[1], "acct": str(r[3]), "cust": r[5], "memo": r[7], "dr": r[8] or 0, "cr": r[9] or 0})

invoices = []
for f in ("07_매출전자세금계산서목록_2025.xlsx", "08_매출전자세금계산서목록_2026_1-2월.xlsx"):
    wb = load_workbook(SRC / f, read_only=True)
    ws = wb["세금계산서"]
    for r in ws.iter_rows(min_row=7, values_only=True):
        if r[1]:
            invoices.append({"appr": r[1], "wdate": r[0], "idate": r[2], "bizno": r[9], "name": r[11], "total": r[14],
                             "supply": r[15], "vat": r[16], "kind": r[17], "note": r[20] or "", "item": r[26]})

bank = [{"dt": r[0], "kind": r[1], "memo": r[2], "out": r[3] or 0, "in": r[4] or 0}
        for r in rows("09_은행거래내역_KB국민_보통예금.xlsx", 7) if r[0]]

ar = [{"code": r[0], "opening": r[4] or 0, "balance": r[7] or 0} for r in rows("05_거래처원장_외상매출금_2025.xlsx", 6) if r[0]]
vat_return = [r[6] for r in rows("10_부가가치세신고서_과세표준요약_2025.xlsx", 4) if r[0]]

files = [
    {"key": "journal", "label": "분개장 (더존 형식)", "names": ["02_분개장_2025.xlsx", "03_분개장_2026_1-2월.xlsx"]},
    {"key": "hometax", "label": "홈택스 매출 전자세금계산서 목록", "names": ["07_매출전자세금계산서목록_2025.xlsx", "08_매출전자세금계산서목록_2026_1-2월.xlsx"]},
    {"key": "bank", "label": "은행 거래내역 (KB국민 보통예금)", "names": ["09_은행거래내역_KB국민_보통예금.xlsx"]},
    {"key": "ar", "label": "거래처원장 · 거래처등록현황", "names": ["05_거래처원장_외상매출금_2025.xlsx", "06_거래처등록현황.xlsx"]},
    {"key": "vat", "label": "부가가치세 신고서 (과세표준 요약)", "names": ["10_부가가치세신고서_과세표준요약_2025.xlsx"]},
]

OUT.mkdir(parents=True, exist_ok=True)
json.dump({"company": "샘플제조(주)", "fiscalYear": 2025, "files": files, "customers": customers, "journal": journal,
           "invoices": invoices, "bank": bank, "ar": ar, "vatReturn": vat_return},
          open(OUT / "input.json", "w"), ensure_ascii=False, separators=(",", ":"))

m = json.load(open(ROOT / "poc" / "output" / "metrics.json"))
json.dump({"anomalies": [{k: a[k] for k in ("id", "type", "assertion", "desc", "amount", "asis", "asis_why", "tobe_no")} for a in m["anomalies"]],
           "coverage": m["coverage"], "asisEffort": m["asis_effort"], "tobeEffort": m["tobe_effort"], "asisVatDiff": m["asis_vat_diff"],
           "asisFindings": m["asis_findings"]},
          open(OUT / "asis.json", "w"), ensure_ascii=False, indent=1)
print("journal", len(journal), "invoices", len(invoices), "bank", len(bank), "customers", len(customers))
