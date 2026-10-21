"""매입 대사 시연 자료: 홈택스 매입 전자세금계산서 목록(가상)을 만든다.

poc의 분개장(외상매입금 251)·은행 출금은 그대로 두고, 그 거래에 대응하는 매입 세금계산서 목록을 만든 뒤
이상거래 6건을 '세금계산서 쪽'에만 심는다. (장부·통장은 회사 자료 그대로)
실행: poc/.venv/bin/python web/scripts/make_purchase_samples.py
출력: web/public/samples/11_매입전자세금계산서목록_2025.xlsx, 12_매입전자세금계산서목록_2026_1-2월.xlsx,
      web/src/data/mock/tieout/purchase_truth.json (심어 둔 이상거래 정답지 — 엔진은 읽지 않음)
"""
import hashlib
import json
import random
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "poc" / "data" / "company"
OUT = ROOT / "web" / "public" / "samples"
rng = random.Random(20260221)

COMPANY = {"bizno": "123-81-45678", "name": "샘플제조(주)", "ceo": "박정민", "addr": "경기도 안산시 단원구 가상산업로 100", "email": "tax@sample-mfg.example"}

purch = []  # 원장 매입 (251 대변)
for f in ("02_분개장_2025.xlsx", "03_분개장_2026_1-2월.xlsx"):
    ws = load_workbook(SRC / f, read_only=True).active
    for r in ws.iter_rows(min_row=6, values_only=True):
        if r[0] and str(r[3]) == "251" and (r[9] or 0) > 0:
            purch.append({"date": datetime.strptime(r[0], "%Y-%m-%d").date(), "memo": r[7], "total": r[9]})
suppliers = sorted({p["memo"].replace("원재료 매입 ", "") for p in purch})
info = {s: {"bizno": f"{rng.randint(200, 699)}-{rng.choice(['81', '86', '87', '88'])}-{rng.randint(10000, 99999)}",
            "ceo": rng.choice(["김도현", "이준호", "박성민", "최영수", "정다은", "한지우"]),
            "addr": f"{rng.choice(['경기도 화성시', '인천광역시 남동구', '충청남도 천안시', '경상남도 창원시'])} 가상로 {rng.randint(10, 300)}"}
        for s in suppliers}

seq = {}


def appr(d):
    k = d.strftime("%Y%m%d")
    seq[k] = seq.get(k, 0) + 1
    h = hashlib.md5(f"{k}{seq[k]}{rng.random()}".encode()).hexdigest()[:8]
    return f"{k}-4200{seq[k]:04d}-{h}"


invoices = []


def inv(s, wdate, total, kind="일반", note="", tag=None):
    supply = round(total / 1.1)
    invoices.append({"wdate": wdate, "appr": appr(wdate), "s": s, "total": total, "supply": supply, "vat": total - supply,
                     "kind": kind, "note": note, "tag": tag})
    return invoices[-1]


truth = []
by_s = {s: [p for p in purch if p["memo"].endswith(s)] for s in suppliers}
S = suppliers  # 6곳

# 정상: 원장 매입 1건 = 세금계산서 1건 (작성일 = 전표일)
skip, shift, mismatch = set(), {}, {}
# P4 원장O·세금계산서X: 5월 매입 세금계산서 없음
p4 = next(p for p in by_s[S[0]] if p["date"].month == 5 and p["date"].year == 2025)
skip.add(id(p4))
truth.append({"id": "P4", "type": "원장O·세금계산서X", "assertion": "발생", "amount": p4["total"],
              "desc": f"{S[0]} 5월 매입 {p4['total']:,}원: 장부·지급은 있으나 국세청 세금계산서 없음 (가공 매입·매입세액 부당공제 위험)"})
# P2 기간귀속: 2026-01 장부 매입의 세금계산서 작성일이 2025-12-29 → 당기 매입·부채 과소
p2 = next(p for p in by_s[S[1]] if p["date"].year == 2026 and p["date"].month == 1)
shift[id(p2)] = date(2025, 12, 29)
truth.append({"id": "P2", "type": "기간귀속(당기 매입 과소)", "assertion": "기간귀속·완전성", "amount": p2["total"],
              "desc": f"{S[1]} 세금계산서 작성일 2025-12-29, 장부는 {p2['date']} 계상 → 2025년 매입·미지급금 누락"})
# P3 금액 불일치: 8월 매입
p3 = next(p for p in by_s[S[2]] if p["date"].month == 8 and p["date"].year == 2025)
mismatch[id(p3)] = p3["total"] - 2_970_000
truth.append({"id": "P3", "type": "금액 불일치", "assertion": "정확성", "amount": 2_970_000,
              "desc": f"{S[2]} 8월 장부 {p3['total']:,}원 vs 세금계산서 {p3['total'] - 2_970_000:,}원 (매입 과대 계상)"})

for p in purch:
    if id(p) in skip:
        continue
    wd = shift.get(id(p), p["date"])
    inv(next(s for s in S if p["memo"].endswith(s)), wd, mismatch.get(id(p), p["total"]))

# P1 미기록 부채: 12월 세금계산서가 국세청엔 있으나 장부·지급 모두 없음
x = inv(S[3], date(2025, 12, 26), 37_400_000, tag="P1")
truth.append({"id": "P1", "type": "세금계산서O·원장X·지급X", "assertion": "완전성", "amount": 37_400_000,
              "desc": f"{S[3]} 2025-12-26 37,400,000원: 장부 매입·미지급금 없음, 2026-02까지 지급도 없음 (미기록 부채)"})
# P5 매입 에누리 미반영: 11월 수정세금계산서(-)가 장부에 없음
o = next(i for i in invoices if i["s"] == S[4] and i["wdate"].month == 10 and i["wdate"].year == 2025)
inv(S[4], date(2025, 11, 14), -3_300_000, kind="수정", note=f"당초승인번호 {o['appr']} 단가 에누리")
truth.append({"id": "P5", "type": "수정세금계산서(매입 차감) 미반영", "assertion": "정확성", "amount": 3_300_000,
              "desc": f"{S[4]} 11월 단가 에누리 -3,300,000원이 장부에 반영 안 됨 (원가·미지급금 과대)"})
# P6 중복 수취: 같은 날·같은 금액 세금계산서가 2건 (장부는 1건)
d = next(i for i in invoices if i["s"] == S[5] and i["wdate"].month == 9 and i["wdate"].year == 2025)
inv(S[5], d["wdate"], d["total"])
truth.append({"id": "P6", "type": "세금계산서 중복 수취", "assertion": "정확성(세무)", "amount": d["total"],
              "desc": f"{S[5]} {d['wdate']} 같은 금액 세금계산서 2건 (장부는 1건) → 매입세액 이중 공제 위험"})

HEAD = ["작성일자", "승인번호", "발급일자", "전송일자", "공급자사업자등록번호", "종사업장번호", "상호", "대표자명", "주소",
        "공급받는자사업자등록번호", "종사업장번호", "상호", "대표자명", "주소", "합계금액", "공급가액", "세액", "전자세금계산서분류",
        "전자세금계산서종류", "발급유형", "비고", "영수/청구 구분", "공급자 이메일", "공급받는자 이메일1", "공급받는자 이메일2",
        "품목일자", "품목명", "품목규격", "품목수량", "품목단가", "품목공급가액", "품목세액", "품목비고"]


def write(path, rows, frm, to):
    wb = Workbook()
    ws = wb.active
    ws.title = "세금계산서"
    ws.append(["사업자 등록번호", COMPANY["bizno"], "상호", COMPANY["name"], "대표자 명", COMPANY["ceo"]])
    ws.append([])
    ws.append(["총 합계금액", sum(r["total"] for r in rows), "총 공급가액", sum(r["supply"] for r in rows), "총 세액", sum(r["vat"] for r in rows)])
    ws.append([])
    ws.append([f"매입 전자(수정)세금계산서 목록조회 (작성일자 {frm} ~ {to})"])
    ws.append(HEAD)
    for r in sorted(rows, key=lambda r: (r["wdate"], r["appr"])):
        i = info[r["s"]]
        sd = (r["wdate"] + timedelta(days=1)).isoformat()
        ws.append([r["wdate"].isoformat(), r["appr"], r["wdate"].isoformat(), sd, i["bizno"], None, r["s"], i["ceo"], i["addr"],
                   COMPANY["bizno"], None, COMPANY["name"], COMPANY["ceo"], COMPANY["addr"], r["total"], r["supply"], r["vat"],
                   "세금계산서" if r["kind"] == "일반" else "수정세금계산서", r["kind"], "정발급", r["note"], "청구",
                   f"tax@{hashlib.md5(r['s'].encode()).hexdigest()[:4]}.supplier.example", COMPANY["email"], None, r["wdate"].isoformat(), "원재료",
                   "", None, None, r["supply"], r["vat"], ""])
    wb.save(path)


write(OUT / "11_매입전자세금계산서목록_2025.xlsx", [r for r in invoices if r["wdate"].year == 2025], "2025-01-01", "2025-12-31")
write(OUT / "12_매입전자세금계산서목록_2026_1-2월.xlsx", [r for r in invoices if r["wdate"].year == 2026], "2026-01-01", "2026-02-28")
json.dump({"note": "엔진은 읽지 않는 정답지. 매입 세금계산서 목록에만 심었다.", "suppliers": suppliers, "anomalies": truth},
          open(ROOT / "web" / "src" / "data" / "mock" / "tieout" / "purchase_truth.json", "w"), ensure_ascii=False, indent=1)
print("suppliers", suppliers, "ledger purchases", len(purch), "invoices", len(invoices))
for t in truth:
    print(" ", t["id"], t["type"], t["amount"])
